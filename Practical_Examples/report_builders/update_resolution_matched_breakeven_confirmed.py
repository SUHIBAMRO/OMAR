"""Updates Table 18-R10e' with the clean, GPU-confirmed re-run
(2026-09-17) of Round6_ResolutionMatchedBreakEven_AllCases.ipynb, after
fixing the memory-cleanup bug (real reference-leak in the except block,
see cell_resolution_matched_break_even_all_cases.py's own comment).
This re-run succeeded cleanly for all four non-Arruda-Boyce cases in
ONE single pass -- no cascade failure -- confirming the code fix
actually works on real GPU, not just reasoned through. Replaces the
asterisked B2xNeo-Hookean substitution (an earlier, separately-measured
clean number used as a workaround) with this run's own number, and
edits the table's own footnote paragraph to say so, in place.

Both the table cell values and the explanatory paragraph are edited
IN PLACE (this document's own established pattern for correcting
already-published content, see add_direct_n1401_ablation_result_to_
report.py), not appended after -- the old asterisk/footnote is simply
superseded, not additional information.
"""
import os

from docx import Document

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-16e.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17.docx')

doc = Document(SRC)

TARGET_HEADER = ['Case', 'torch-fem @N=1401', 'Operator @N=1401 (eager fp32)', 'Speedup', 'Break-even']
NEW_ROWS = {
    'B1 x Neo-Hookean': ['135.00 s', '2292.1 ms', '58.9x', '316 samples'],
    'B1 x Mooney-Rivlin': ['133.92 s', '2361.9 ms', '56.7x', 'training cost unknown'],
    'B2 x Neo-Hookean': ['205.14 s', '2351.1 ms', '87.3x', 'training cost unknown'],
    'B2 x Mooney-Rivlin': ['207.21 s', '2356.0 ms', '88.0x', 'training cost unknown'],
}

found_table = False
for t in doc.tables:
    header = [c.text for c in t.rows[0].cells]
    if header != TARGET_HEADER:
        continue
    found_table = True
    for row in t.rows[1:]:
        case = row.cells[0].text
        if case in NEW_ROWS:
            new_vals = NEW_ROWS[case]
            for j, v in enumerate(new_vals, start=1):
                row.cells[j].text = v
    break
assert found_table, "Table 18-R10e' not found -- check the header text matches exactly"


def replace_paragraph_text(para, new_text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(new_text)


OLD_FOOTNOTE = (
    "Table 18-R10e'. Resolution-matched break-even: operator vs. torch-fem, both "
    "evaluated at N=1401, all six cases. At this matched resolution the operator "
    "wins by 57x-89x even in its default, unoptimized deployment mode -- markedly "
    "more favourable than the accuracy-matched comparison above (Table 18-R10e), "
    "where the same default mode does not break even at all and the operator's "
    "economic case instead rests on the torch.compile+TF32 optimization from point "
    "3. \"Break-even (samples)\" is only reported where that case's own retrained-"
    "checkpoint training wall-clock is known; for the three cases marked \"training "
    "cost unknown\" only the per-sample speedup is available, since those "
    "checkpoints predate this project's own training-time logging. * B2 x "
    "Neo-Hookean's row uses an earlier, independently clean measurement rather than "
    "the same session's freshest re-run of this sweep: that re-run's own attempt at "
    "this case failed with a CUDA out-of-memory error, but the sweep's own printed "
    "GPU-memory state shows 81.27 GB already allocated on a 79.25 GB device before "
    "this case's forward pass even started -- memory left over from the immediately "
    "preceding B1 x Arruda-Boyce failure, never released before the next case began. "
    "That is a memory-cleanup gap between cases in the sweep script, not a real "
    "finding about this case, which is why the earlier clean number is used here "
    "instead of silently reporting the corrupted re-run's own failure as if it were "
    "a genuine result."
)
NEW_FOOTNOTE = (
    "Table 18-R10e'. Resolution-matched break-even: operator vs. torch-fem, both "
    "evaluated at N=1401, all six cases. At this matched resolution the operator "
    "wins by 57x-89x even in its default, unoptimized deployment mode -- markedly "
    "more favourable than the accuracy-matched comparison above (Table 18-R10e), "
    "where the same default mode does not break even at all and the operator's "
    "economic case instead rests on the torch.compile+TF32 optimization from point "
    "3. \"Break-even (samples)\" is only reported where that case's own retrained-"
    "checkpoint training wall-clock is known; for the three cases marked \"training "
    "cost unknown\" only the per-sample speedup is available, since those "
    "checkpoints predate this project's own training-time logging. All four "
    "non-Arruda-Boyce numbers above are from a single clean re-run (2026-09-17), "
    "after fixing a real memory-cleanup bug in the sweep script found on a previous "
    "attempt (a reference to the OOM exception object, independent of the "
    "auto-deleted `except ... as e` binding, kept a failed case's own GPU tensors "
    "pinned in memory into the next case). This re-run succeeded for every "
    "non-Arruda-Boyce case in one pass with no cascade failure, confirming the fix "
    "on real GPU rather than only by code inspection."
)

replaced = False
for p in doc.paragraphs:
    if p.text.strip() == OLD_FOOTNOTE:
        replace_paragraph_text(p, NEW_FOOTNOTE)
        replaced = True
        break
assert replaced, "footnote paragraph not found -- check the text matches exactly"

doc.save(DST)
print('Saved', DST)
