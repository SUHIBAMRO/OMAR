"""Adds the resolution-matched break-even speedup figure (drafted and
approved by Omar first, see fig_resolution_matched_breakeven.png) right
after Table 18-R10e''s own footnote paragraph -- the natural place for
a figure illustrating the same table."""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

FIG = '/home/user/OMAR/Practical_Examples/report_builders/figures'
DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17b.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_figure_after(anchor_para, image_path, caption_text, width_in=6.0):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width_in))
    img_elem = img_p._p
    img_elem.getparent().remove(img_elem)

    cap_p = doc.add_paragraph()
    cap_run = cap_p.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_elem = cap_p._p
    cap_elem.getparent().remove(cap_elem)

    anchor_para._p.addnext(cap_elem)
    anchor_para._p.addnext(img_elem)
    return Paragraph(cap_elem, anchor_para._parent)


anchor = find_para_exact(
    "Table 18-R10e'. Resolution-matched break-even: operator vs. torch-fem, both "
    "evaluated at N=1401, all six cases. At this matched resolution the operator "
    "wins by 57x-89x even in its default, unoptimized deployment mode -- markedly "
    "more favourable than the accuracy-matched comparison above (Table 18-R10e), "
    "where the same default mode does not break even at all and the operator's "
    "economic case instead rests on the torch.compile+TF32 optimization from point "
    "3. \"Break-even (samples)\" is only reported where that case's own retrained-"
    "checkpoint training wall-clock is known; for the three cases marked \"training "
    "cost unknown\" only the per-sample speedup is available, since those "
    "checkpoints predate this project's own training-time logging. All four "
    "non-Arruda-Boyce numbers above are from a single clean re-run (2026-09-17), "
    "after fixing a real memory-cleanup bug in the sweep script found on a previous "
    "attempt (a reference to the OOM exception object, independent of the "
    "auto-deleted `except ... as e` binding, kept a failed case's own GPU tensors "
    "pinned in memory into the next case). This re-run succeeded for every "
    "non-Arruda-Boyce case in one pass with no cascade failure, confirming the fix "
    "on real GPU rather than only by code inspection."
)

insert_figure_after(anchor,
    os.path.join(FIG, 'fig_resolution_matched_breakeven.png'),
    "Round-10 Figure E. Resolution-matched speedup (operator vs. torch-fem, both "
    "at N=1401), all six cases (Table 18-R10e'). Both Arruda-Boyce cases shown as "
    "FAILED rather than omitted, since torch-fem's own solve does not converge for "
    "either at this resolution.")

doc.save(DST)
print('Saved', DST)
