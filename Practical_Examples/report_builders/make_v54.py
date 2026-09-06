"""v53 -> v54. Timon's round-8 review (2026-09-08), point 4: "Can you
please provide also the wall clock training time for DD-NO." The
numbers were already committed in
point7b_results/comparison_B1_neo_hookean.json (they just were not
placed in Table 21) -- adds two rows for training wall-clock, matching
the caption's own "All four runs used 75,000 optimiser steps at batch
size 8" so all four training-time values sit next to the error values
they belong with.
"""
import json
import os

from docx import Document

SRC, DST = 'PFEM_Transolver_Report_v53.docx', 'PFEM_Transolver_Report_v54.docx'
HERE = os.path.dirname(os.path.abspath(__file__))
PF = os.path.join(HERE, '..', 'omar_pfem')

D = json.load(open(os.path.join(PF, 'point7b_results', 'comparison_B1_neo_hookean.json')))
runs = D['runs']

pi_adam = runs['physics_informed']['train_wall_clock_s']
pi_onecycle = runs['physics_informed_adamw_onecycle']['train_wall_clock_s']
dd_adam = runs['data_driven_matched_optimizer']['train_wall_clock_s']
dd_onecycle = runs['data_driven_own_optimizer']['train_wall_clock_s']

assert (pi_adam, pi_onecycle, dd_adam, dd_onecycle) == (2873.8, 3108.9, 1458.3, 1463.0)

doc = Document(SRC)
t = doc.tables[40]
assert [c.text for c in t.rows[0].cells] == [
    'Training loss', 'Adam, lr 2×10⁻³', 'AdamW lr 10⁻³ + OneCycleLR']
assert t.rows[1].cells[0].text == 'Physics-informed (energy)'
assert t.rows[2].cells[0].text == 'Data-driven (relative L2 to FEM)'

import copy
from docx.table import _Row

for label, adam_val, onecycle_val in [
    ('Physics-informed, wall-clock (s)', pi_adam, pi_onecycle),
    ('Data-driven, wall-clock (s)', dd_adam, dd_onecycle),
]:
    new_tr = copy.deepcopy(t.rows[2]._tr)
    t._tbl.append(new_tr)
    row = _Row(new_tr, t)
    row.cells[0].text = label
    row.cells[1].text = f'{adam_val:,.1f}'
    row.cells[2].text = f'{onecycle_val:,.1f}'

assert len(t.rows) == 5
print('Table 21 now:')
for r in t.rows:
    print([c.text for c in r.cells])

ORIGINAL = list(doc.paragraphs)


def find_para(prefix):
    hits = [p for p in ORIGINAL if p.text.strip().startswith(prefix)]
    assert len(hits) == 1, f'{len(hits)} paragraphs start with {prefix!r}'
    return hits[0]


def retext(prefix, text):
    p = find_para(prefix)
    keep = p.runs[0] if p.runs else p.add_run()
    keep.text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


retext('Table 21. Held-out relative L2 error', (
    'Table 21. Held-out relative L2 error, B1 × Neo-Hookean, for each '
    'combination of training loss and optimiser, plus each run’s '
    'training wall-clock time. Reading down a column isolates the loss, '
    'because everything else in that column is identical. All four runs '
    'used 75,000 optimiser steps at batch size 8. The data-driven runs’ '
    'wall-clock excludes the 800 FEM solves used to generate their '
    'labels (5.65 h of CPU time, paid once, not part of training itself).'
))

doc.save(DST)
print('wrote', DST)
