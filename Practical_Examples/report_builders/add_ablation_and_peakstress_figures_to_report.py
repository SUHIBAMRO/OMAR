"""Adds the two draft figures (approved by Omar after two color
iterations) to the Report: one for Table 18-R10j (direct-N1401
ablation) and one for Table 18-R10k (peak-stress QoI, 5 remaining
cases). Each is inserted right after its own table's caption paragraph
and before the interpretation prose that follows it -- the same
Table -> caption -> Figure -> interpretation order already established
for Round-10 Figure A (Table 18-R10b) elsewhere in this document.
"""
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

FIG = '/home/user/OMAR/Practical_Examples/report_builders/figures'
DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17b.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17c.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_figure_after(anchor_para, image_path, caption_text, width_in=6.0):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width_in))
    img_elem = img_p._p
    img_elem.getparent().remove(img_elem)

    cap_p = doc.add_paragraph()
    cap_run = cap_p.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_elem = cap_p._p
    cap_elem.getparent().remove(cap_elem)

    anchor_para._p.addnext(cap_elem)
    anchor_para._p.addnext(img_elem)
    return Paragraph(cap_elem, anchor_para._parent)


anchor_ablation = find_para_exact(
    "Table 18-R10j. Direct-N1401 training ablation vs. the multi-resolution "
    "zero-shot checkpoint, same architecture, same material (B1 x Neo-Hookean)."
)
insert_figure_after(anchor_ablation,
    os.path.join(FIG, 'fig_direct_n1401_ablation.png'),
    "Round-10 Figure F. Direct-N1401 training vs. multi-resolution zero-shot, "
    "training cost and accuracy (Table 18-R10j).")

# Re-fetch: the ablation insertion shifted every later paragraph's index, but
# find_para_exact re-scans ORIGINAL (captured before any insertion), which is
# still valid since ORIGINAL's own paragraph objects are unaffected by later
# insertions elsewhere in the tree -- only their position, not their identity.
anchor_peak = find_para_exact(
    "Table 18-R10k. Peak-stress QoI at N=1401, the five cases beyond "
    "B1xNeo-Hookean (already covered above). Same fixed-location, fine-reference "
    "convention as the B1xNeo-Hookean peak-stress discussion."
)
insert_figure_after(anchor_peak,
    os.path.join(FIG, 'fig_peak_stress_qoi.png'),
    "Round-10 Figure G. Peak-stress QoI vs. displacement error, five cases "
    "(Table 18-R10k).")

doc.save(DST)
print('Saved', DST)
