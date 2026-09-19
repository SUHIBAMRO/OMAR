"""Round-13 non-cumulative Summary (Omar's own 2026-09-18 convention:
Summary covers ONLY the round just completed, copied verbatim out of the
now-updated cumulative Report -- see build_new_summary_2026-09-18.py for
the original rationale/mechanism, reused unchanged here).

Round 13 = Timon's newest email ("let's wrap up the benchmark work"),
items 1 and 2 (item 3, dropping IGA/NURBS, needs no report content --
explicit future work only; item 4, the new 3D example, is its own much
larger deliverable and will get its own round/Summary later).
"""
import copy
import os

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

DELIV = '/home/user/OMAR/advisor_feedback'
REPORT = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-19b.docx')
OUT = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-19.docx')

src = Document(REPORT)


def body_children(doc):
    return list(doc.element.body)


def elem_text(el):
    if el.tag != qn('w:p'):
        return None
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_index(children, text):
    matches = [i for i, el in enumerate(children)
               if elem_text(el) is not None and elem_text(el).strip() == text]
    assert len(matches) == 1, f'{len(matches)} matches for {text!r}'
    return matches[0]


def slice_by_anchor_text(doc, start_text, end_text):
    children = body_children(doc)
    i0 = find_index(children, start_text)
    i1 = find_index(children, end_text)
    assert i1 >= i0
    if i1 + 1 < len(children) and children[i1 + 1].tag == qn('w:tbl'):
        i1 += 1
    return children[i0:i1 + 1]


def append_heading(dst_doc, text):
    from docx.shared import Pt
    p = dst_doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)


def append_quote(dst_doc, text):
    from docx.shared import Pt
    p = dst_doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)


def copy_elements(dst_doc, src_doc, elements):
    body = dst_doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    for el in elements:
        new_el = copy.deepcopy(el)
        for blip in new_el.iter(qn('a:blip')):
            rId = blip.get(qn('r:embed'))
            if not rId:
                continue
            image_part = src_doc.part.related_parts[rId]
            new_rId = dst_doc.part.relate_to(image_part, RT.IMAGE)
            blip.set(qn('r:embed'), new_rId)
        if sect_pr is not None:
            sect_pr.addprevious(new_el)
        else:
            body.append(new_el)


dst = Document()
if src.tables:
    src_style = src.tables[0].style
    try:
        dst.styles.add_style(src_style.name, src_style.type)
    except ValueError:
        pass

append_heading(dst, "PFEM / Transolver Work Summary -- Round 13")
p = dst.add_paragraph()
p.add_run(
    "This round covers items 1 and 2 of the advisor's follow-up email after round 12 "
    "(\"let's wrap up the benchmark work\"). Item 3 (drop IGA/NURBS) needed no report "
    "content -- confirmed explicit future work only, discussed after this paper. Item 4 "
    "(a new, harder 3D example) is its own much larger deliverable and will get its own "
    "round and Summary once built, per the advisor's own stated order. As with round 12, "
    "everything below is copied verbatim (same tables, same text) out of the now-updated "
    "cumulative Report (PFEM_Transolver_Report_2026-09-19b.docx) -- nothing summarized or "
    "re-derived."
).italic = True

# =======================================================================
# ITEM 1 -- QoI-accuracy-threshold vs. required-FEM-N vs. break-even
# =======================================================================
append_heading(dst, "Item 1 -- QoI-accuracy-threshold vs. required-FEM-resolution vs. break-even")
append_quote(dst,
    "“Finalize the low-N accuracy comparison (drop N=1401 for B1/B2) and add a table "
    "showing, for a few QoI-accuracy thresholds, the required FEM resolution and the "
    "break-even point.” — Timon"
)
elements = slice_by_anchor_text(
    src,
    "Round-13 (the advisor's follow-up email after round 12): finalize the low-N "
    "comparison and add one new analysis -- for a few QoI-accuracy thresholds (1%, 2%, "
    "5%, the advisor's own examples), per QoI, the minimum GPU-native-FEM resolution N "
    "required to reach it, and the operator's own break-even point there. This reuses "
    "round 12's own consistent-field accuracy data (FEM and the operator on the "
    "identical field, scored against one real fine reference) for the accuracy side, "
    "and adds one new real measurement for the cost side: our own GPU-native FEM "
    "solver's (gpu_fem_solver.py) per-sample wall-clock at every one of the sixteen "
    "LOW_N resolutions, not just the single N=11 point the earlier accuracy-matched "
    "break-even table (Table 18-R10e) used. Break-even always uses the operator's own "
    "compile+TF32 cost at N=1401 (its real deployment resolution, essentially "
    "case-independent at ~394ms -- see Table 18-R10h) against that resolution's own "
    "training wall-clock -- the SAME operator baseline as the existing break-even "
    "tables, not a new one. One table per case below.",
    "Two findings worth stating plainly. First, across all six cases and every "
    "threshold tested, the required FEM resolution never exceeds N=41, and its "
    "own per-sample cost there (up to ~7.1s) never comes close to the operator's "
    "own ~394ms -- FEM never becomes cheaper than the operator at any accuracy level "
    "tested, so a break-even point always exists and is always finite; the six cases' "
    "break-even sample counts differ almost entirely by each case's own training cost, "
    "not by its accuracy numbers (compare B1xMooney-Rivlin's ~34,500 samples, driven by "
    "its 14.81h training run, against B2xMooney-Rivlin's ~4,000 samples, driven by its "
    "own much cheaper 1.76h run, at the SAME L2@2% threshold). Second, H1 semi-norm and "
    "tangent energy never reach the 1% threshold anywhere in the tested range (N=3..49) "
    "for any of the six cases -- consistent with this project's own repeated finding "
    "that energy/H1-type norms converge more slowly than displacement L2, and a real, "
    "disclosed gap: a 1% energy/H1 target would need FEM resolutions finer than this "
    "sweep covers, not a number this table can currently supply.",
)
copy_elements(dst, src, elements)

# =======================================================================
# ITEM 2 -- name gpu_fem_solver.py explicitly as the primary GPU-native baseline
# =======================================================================
append_heading(dst, "Item 2 -- confirm the GPU-native solver as the primary timing baseline")
append_quote(dst,
    "“Use your most efficient validated GPU-native FEM solver as the primary timing "
    "baseline; keep torch-fem/TensorMesh as secondary.” — Timon"
)
p = dst.add_paragraph()
p.add_run(
    "Investigated via direct code inspection, not assumption: Table 18-R10e's own "
    "finite-element measurement already comes from gpu_fem_benchmark.py -> "
    "gpu_fem_solver.py -- our own GPU-native Total-Lagrangian Newton solver, built "
    "earlier per the advisor's own request for a GPU-native comparison, completely "
    "independent of torch-fem. Torch-fem is used only in the separate resolution-matched "
    "table (18-R10e'), matching the advisor's own request to keep that comparison apart. "
    "So this item needed no computation change -- only making it explicit in the Report "
    "text, since the cell producing this number only named it \"torch-fem\" in its own "
    "informal print() statements, never in the actual Report wording. Two paragraphs "
    "(below) were updated accordingly; nothing else changed."
).italic = True
elements = slice_by_anchor_text(
    src,
    "Point 5 asked for an updated break-even analysis. Using the retrained checkpoint's "
    "own coarsest-suitable finite-element mesh at N=1401 (N=11, bound by the "
    "tangent-energy norm per the crossover above) against the operator's own real "
    "inference cost there, and the real training wall-clock for the retrained checkpoint "
    "(41,881 s, about 11.6 hours, confirmed above): in its default eager fp32 deployment "
    "mode, the operator never recovers its own training cost against this baseline -- the "
    "accuracy-matched finite-element mesh is already cheaper per sample (1,625.6 ms) than "
    "the operator's default forward pass (2,292.1 ms). With the torch.compile plus TF32 "
    "optimization above (394.0 ms/sample), the operator is instead 4.13x faster per "
    "sample and recovers its training cost after 34,005 solved problems, about 3.72 "
    "GPU-hours of inference against the 11.6 GPU-hours spent training. This is the only "
    "point in the whole round-10 investigation where the operator's default, unoptimized "
    "deployment mode does not come out ahead; its economic case against a genuinely "
    "accuracy-matched baseline rests on deploying it with the optimizations found under "
    "point 3, not on its default settings. (The finite-element side of this comparison is "
    "our own GPU-native Total-Lagrangian Newton solver, gpu_fem_solver.py -- built earlier "
    "per the advisor's own request for a GPU-native comparison, and already the primary "
    "timing baseline used here; torch-fem appears only in the separate resolution-matched "
    "comparison below, Table 18-R10e', which intentionally asks a different question -- "
    "same-N, not accuracy-matched -- per the advisor's own request to keep the two "
    "comparisons apart.)",
    "Table 18-R10e. Accuracy-matched break-even: operator@N=1401 vs. our own GPU-native "
    "finite-element solver (gpu_fem_solver.py) @N=11 (its coarsest suitable mesh for the "
    "retrained checkpoint). Distinct from the matched-resolution/matched-batch-size "
    "break-even of Section 8.4: this comparison intentionally runs the two methods at "
    "different N, matched by accuracy rather than by mesh.",
)
copy_elements(dst, src, elements)

dst.save(OUT)
print('Saved', OUT)

check = Document(OUT)
print('paragraphs:', len(check.paragraphs))
print('tables:', len(check.tables))
print('images:', len(check.inline_shapes))
