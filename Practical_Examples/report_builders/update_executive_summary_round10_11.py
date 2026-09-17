"""Omar caught this during the pre-send audit: the Report's own
Executive Summary still says "This revision addresses all seven points
you raised in your last feedback" -- stale from round-9, never updated
to mention any of round-10 or round-11's own work, even though that
material has been in the document body for a while now (checkpoint bug
fix, multi-resolution retrain, N=1401 analysis extended to all six
cases, direct-N1401 ablation, both break-even comparisons). Adds one
concise paragraph recapping round-10/11 at the same level of brevity as
the existing round-9 sentence, and states the one item still open
(point 4, complex geometry) exactly as the body itself already frames
it.

Pure addition -- the existing round-9 sentence and scope note are left
untouched, since round-9's own claims are still true and this is a
Report (unlike the Summary, which Omar asked to trim to round-10/11
only), not a document being cut down to one round's scope.
"""
import copy
import os

from docx import Document
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17f.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17g.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


anchor = find_para_exact(
    "Scope note: one item is not extended to all six benchmark cases. The "
    "~10-million/40-million-DOF Q4-vs-Q9 convergence study of Section 4.4 "
    "(point 1's deeper, numerical-reference-based check, distinct from the "
    "h-refinement sweep of Section 4.3, which is confirmed for all six) is "
    "deliberately confined to B1; see Section 4.4. Every other point above, "
    "including the resolution-invariance study of point 7, is confirmed "
    "across all six (geometry, material) combinations (Section 10)."
)

insert_after(
    anchor,
    "This revision was subsequently extended twice further, per two more "
    "rounds of feedback (round-10 and round-11). A real checkpoint-loading "
    "bug was found and fixed (round-10's own first accuracy figures had "
    "silently used the wrong trained model). The operator's accuracy at the "
    "deployment resolution N=1401 was then checked directly against real "
    "finite-element ground truth for the first time and found to degrade "
    "substantially, motivating a multi-resolution retraining fix, validated "
    "across every applicable case, that reduces the error at N=1401 by "
    "roughly 22-87% depending on the case. The N=1401 accuracy/QoI/break-even "
    "analysis was extended to all six geometry-material combinations, "
    "together with an explicit table identifying which quantity of interest "
    "actually determines each case's own coarsest-suitable finite-element "
    "mesh. A direct-N1401 training ablation was run as a control against the "
    "zero-shot multi-resolution result, and both an accuracy-matched and a "
    "resolution-matched break-even comparison are now reported side by side. "
    "One item remains open: the complex-geometry example (round-10/11 point "
    "4), designed and verified with real mesh-convergence evidence, awaiting "
    "your confirmation before any training time is committed to it."
)

doc.save(DST)
print('Saved', DST)
