"""v54 -> v55. Timon's round-8 point 7 asked that the single-query,
matched-batch-size comparison (his "Comparison A": FEM bs=1 vs. NO
inference bs=1) be the PRIMARY comparison, with the batched throughput
figures shown separately rather than as the headline. The report's own
prose already reaches exactly this conclusion (the paragraph after
Table 10d already says the bs=1 column "is the figure that matters for
a deployment claim" and that the bs=128 column is "the wrong one to
quote as the headline") -- this just labels it explicitly in his own
terms, so there is no ambiguity when he rereads it.
"""
from docx import Document

SRC, DST = 'PFEM_Transolver_Report_v54.docx', 'PFEM_Transolver_Report_v55.docx'

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para(prefix):
    hits = [p for p in ORIGINAL if p.text.strip().startswith(prefix)]
    assert len(hits) == 1, f'{len(hits)} paragraphs start with {prefix!r}'
    return hits[0]


def insert_after(prefix, text):
    import copy
    from docx.text.paragraph import Paragraph
    p = find_para(prefix)
    new_p = copy.deepcopy(p._p)
    p._p.addnext(new_p)
    np = Paragraph(new_p, p._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


insert_after(
    'What this clarifies is where the operator is useful',
    'To name this plainly: the batch-size-1 comparison above is the '
    'primary, single-query result — matched hardware, matched batch '
    'size, one FEM solve against one operator inference — and every '
    'other batch size in Tables 10 through 10d is a separate throughput '
    'experiment, useful for understanding how both methods scale with '
    'batching but not the figure a deployment claim should be built on.'
)

doc.save(DST)
print('wrote', DST)
