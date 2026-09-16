"""Adds one clarifying sentence to the MMS section (Gap 3 from the
2026-09-16 audit) answering a question Timon raised directly: since
every manufactured solution generates its own body-force field
b = -Div P, how is that varying, sample-dependent forcing communicated
to the separate MMS operator -- through the family's (alpha, beta)
parameters, or directly as an input field? Investigated in
mms_operator.py this session: the operator receives the SAME assembled
nodal force vector FEM's own right-hand side uses (both call
assemble_body_force with identical arguments), fed directly as a
per-node input channel, only statistically normalized -- not encoded
through alpha/beta, which are used only to GENERATE the manufactured
family, never seen by the operator itself. This directly supports the
"fairness" comparison Timon asked about: the operator is not required
to infer the varying source implicitly, it receives the same
discretized load FEM does.

Inserted immediately after the paragraph introducing the separate MMS
operator ("A separate operator was therefore trained for this
study..."), the natural place for this clarification."""
import copy
import os

from docx import Document
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-16d.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-16e.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


anchor = find_para_exact(
    "The third leg. The comparison the advisor asked for is three-way, and the two "
    "tables above are FEM only. The report's trained operators cannot be pointed "
    "at a manufactured problem: their energy functional has no body-force term and "
    "their input channels have no field to carry one, so the checkpoints of Table "
    "5 do not apply. A separate operator was therefore trained for this study the "
    "same architecture, on the same Q4 mesh, minimising the same discrete "
    "potential energy, and scored by the same error routine as the two solvers, so "
    "the three columns are commensurable even though the third model is not one of "
    "the report's own."
)

insert_after(anchor,
    "One fairness question the advisor raised directly deserves stating "
    "explicitly here: since every manufactured solution generates its own "
    "body-force field b = -Div P, how does that sample-dependent forcing reach "
    "the operator -- through the manufactured family's own (alpha, beta) "
    "amplitude parameters, or as a direct input field? Checked in the code rather "
    "than assumed: the operator's forcing channel and the FEM solver's own "
    "right-hand side are built by the SAME assembly routine, called with "
    "identical arguments, so they are the identical discretized load, not "
    "independently-derived approximations of it. That assembled nodal force "
    "vector is fed to the operator directly as a per-node input channel (only "
    "statistically normalized before entering the network); (alpha, beta) are "
    "used solely to GENERATE the manufactured family and its forcing, and are "
    "never themselves seen by the operator. The operator is therefore not asked "
    "to infer the varying source implicitly from a low-dimensional parameter -- it "
    "receives the same discretized load the finite-element solver does, which is "
    "the property that makes the three-way comparison here a fair one in this "
    "specific sense.")

doc.save(DST)
print('Saved', DST)
