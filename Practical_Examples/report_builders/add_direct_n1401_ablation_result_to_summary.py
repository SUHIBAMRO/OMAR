"""Companion to add_direct_n1401_ablation_result_to_report.py -- same
in-place edit + new result, condensed to the Summary's own terser
style. See that script's own docstring for full rationale and exact
number provenance."""
import copy
import os

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-16b.docx')
DST = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-16c.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def replace_paragraph_text(para, new_text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(new_text)
    return para


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_table_after(anchor_para, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    pr = doc.tables[0]._tbl.find(qn('w:tblPr'))
    if pr is not None:
        old = tbl._tbl.find(qn('w:tblPr'))
        if old is not None:
            tbl._tbl.remove(old)
        tbl._tbl.insert(0, copy.deepcopy(pr))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    tbl._tbl.getparent().remove(tbl._tbl)
    anchor_para._p.addnext(tbl._tbl)
    return tbl


def insert_after_table(table, style_para, text):
    new_p_el = copy.deepcopy(style_para._p)
    table._tbl.addnext(new_p_el)
    np = Paragraph(new_p_el, style_para._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


anchor = find_para_exact(
    "Results, N=1401, old vs. new checkpoint (Table R10-2): B1 x Mooney-Rivlin "
    "39.20% -> 15.04% (~62% relative reduction); B1 x Arruda-Boyce 45.62% -> 25.05% "
    "(~45%, and its own real GPU run is reassuring evidence that a shared "
    "chain-locking-clamp exposure with this project's own torch-fem Arruda-Boyce "
    "runs does not bite in practice here -- all 16 resolutions, both checkpoints, "
    "converged cleanly); B2 x Mooney-Rivlin 49.47% -> 13.10% (~73.5%, the largest "
    "reduction so far, and the clearest case of a genuinely FLAT error across "
    "N=37-1401 rather than merely a smaller one at N=1401). B2 x Neo-Hookean's own "
    "retrain and the advisor's round-11 direct-N1401 ablation were both still "
    "mid-training as of this writing and are not included below."
)

replace_paragraph_text(anchor,
    "Results, N=1401, old vs. new checkpoint (Table R10-2): B1 x Mooney-Rivlin "
    "39.20% -> 15.04% (~62% relative reduction); B1 x Arruda-Boyce 45.62% -> 25.05% "
    "(~45%, and its own real GPU run is reassuring evidence that a shared "
    "chain-locking-clamp exposure with this project's own torch-fem Arruda-Boyce "
    "runs does not bite in practice here -- all 16 resolutions, both checkpoints, "
    "converged cleanly); B2 x Mooney-Rivlin 49.47% -> 13.10% (~73.5%, the largest "
    "reduction so far, and the clearest case of a genuinely FLAT error across "
    "N=37-1401 rather than merely a smaller one at N=1401). B2 x Neo-Hookean's own "
    "retrain was still mid-training as of this writing and is not included below.")

p1 = insert_after(anchor,
    "Round-11 point 4 -- direct-N1401 training ablation -- has since completed. A "
    "B1 x Neo-Hookean checkpoint trained DIRECTLY at N=1401 (100 samples, early "
    "stopped at epoch 36, best epoch 20) is cheaper to train (8.00h vs. 11.63h, "
    "~30% less GPU time) but 6.3x LESS accurate (36.65% vs. 5.85% disp_rel_L2) than "
    "the multi-resolution zero-shot checkpoint, with essentially identical "
    "inference cost (2318.8 ms vs. 2292.1 ms/sample). A genuine result in favour of "
    "the multi-resolution strategy: not just \"it also works zero-shot,\" but "
    "\"it produces a substantially better model for less than 50% more training "
    "time.\" Table R10-5.")

header = ['Metric', 'Direct @N=1401', 'Multi-res (zero-shot)']
t = insert_table_after(p1, header, [
    ['Training wall-clock', '8.00 h', '11.63 h'],
    ['disp_rel_L2 @ N=1401', '36.65%', '5.85%'],
    ['Inference (eager fp32)', '2318.8 ms', '2292.1 ms'],
])
p2 = insert_after_table(t, p1,
    "Table R10-5. Direct-N1401 training ablation vs. multi-resolution zero-shot, "
    "B1 x Neo-Hookean.")

doc.save(DST)
print('Saved', DST)
