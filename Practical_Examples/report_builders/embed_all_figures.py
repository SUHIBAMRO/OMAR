"""Item #12: embed all 13 table-based figures into the real Report and
Summary docx files, right after the table(s) each one visualizes.

Anchor resolution mirrors docx_table_map.py's own before/after-caption
logic (re-derived per table here since build_table_map() doesn't expose
which side it used for a given table): the true "end" of a table's own
block is whichever of {the table, its caption paragraph} comes LAST in
document order -- if the caption sits after the table (the common case
in this document), that caption paragraph is the anchor; if the caption
sits before it (the six mesh-convergence tables' own reversed
convention), the table itself is the anchor. For a group of tables
(e.g. the six 1/1a/1b/2/2a/2b tables sharing one summary figure), the
group's anchor is whichever single table's own end-anchor sits latest
in the document -- inserting after all of them, not just the first.

IMPORTANT, found by checking the actual result before trusting it: the
Report and Summary do NOT lay out sections in strictly ascending table-
number order (e.g. the MMS section, Tables 22-24, physically sits
BEFORE the B2 fix-history section, Tables 13/14, in this document), and
the two documents don't even order sections the same way as each other.
Assigning "Figure 17, 18, 19..." by table number instead of by each
figure's own actual anchor position would have produced a document
where e.g. "Figure 28" appears on the page before "Figure 22" -- so
figure numbers are computed HERE, per document, from each anchor's real
position, not assumed from the table numbering.

New figures follow the same visual convention already used by this
document's existing 16 embedded figures (confirmed by inspection
before writing this): a centered image paragraph, immediately followed
by an italic 9pt "Figure N. <caption>" paragraph -- nothing else about
the surrounding text is touched.
"""
import os
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(__file__))
from docx_table_map import build_table_map, _nearest_caption

DELIV = '/tmp/claude-0/-home-user/64d7c4d8-d5f0-5686-a58f-aa87abfd4ba4/scratchpad/deliverables'
FIG = os.path.join(os.path.dirname(__file__), 'figures')
REPORT = os.path.join(DELIV, 'PFEM_Transolver_Report_updated_2026-09-08.docx')
SUMMARY = os.path.join(DELIV, 'PFEM_Work_Summary_updated_2026-09-08.docx')
START_FIGURE_NUMBER = 17  # both documents currently have Figures 1-16


def load_items(doc):
    body = doc.element.body
    items = []
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            items.append(('p', Paragraph(child, doc)))
        elif child.tag == qn('w:tbl'):
            items.append(('tbl', Table(child, doc)))
    return items


def table_index(items, table):
    for idx, (kind, obj) in enumerate(items):
        if kind == 'tbl' and obj._tbl is table._tbl:
            return idx
    raise ValueError('table not found among items')


def end_anchor_for_label(items, tmap, label):
    """Returns the anchor XML element for one table label -- the table
    itself, or its caption paragraph, whichever comes later in
    document order."""
    table = tmap[label]
    idx = table_index(items, table)
    after = _nearest_caption(items, idx, +1)
    before = _nearest_caption(items, idx, -1)
    if after == label:
        j = idx + 1
        while items[j][0] == 'p' and not items[j][1].text.strip():
            j += 1
        return items[j][1]._p, j
    elif before == label:
        return table._tbl, idx
    else:
        return table._tbl, idx


def group_anchor(items, tmap, labels):
    """Anchor (element, index) for a group of tables: the end-anchor of
    whichever member table is physically LAST in the document."""
    labels = [l for l in labels if l in tmap]
    assert labels, f'none of the requested labels exist in this document: {labels}'
    best = max(labels, key=lambda l: table_index(items, tmap[l]))
    return end_anchor_for_label(items, tmap, best)


def text_anchor(items, needle):
    for idx, (kind, obj) in enumerate(items):
        if kind == 'p' and needle in obj.text:
            return obj._p, idx
    raise ValueError(f'anchor text not found: {needle!r}')


def insert_figure(anchor_elem, doc, image_path, caption_text, width_in=6.5):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width_in))
    img_elem = img_p._p
    img_elem.getparent().remove(img_elem)

    cap_p = doc.add_paragraph()
    cap_run = cap_p.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_elem = cap_p._p
    cap_elem.getparent().remove(cap_elem)

    anchor_elem.addnext(cap_elem)
    anchor_elem.addnext(img_elem)


# (figure file, base caption (no "Figure N." prefix), report anchor labels,
# summary anchor labels [optional, defaults to report labels])
FIGURES = [
    ('fig_mesh_convergence.png',
     'Mesh convergence: strain-energy change per refinement vs. N, B1 and B2, '
     'all three materials (Tables 1/1a/1b/2/2a/2b).',
     ['Table 1', 'Table 1a', 'Table 1b', 'Table 2', 'Table 2a', 'Table 2b']),
    ('fig_training_cost.png',
     'Training cost, per-sample inference speed-up, and training memory '
     'footprint across all six (geometry, material) cases (Tables 5/7/8).',
     ['Table 5', 'Table 7', 'Table 8']),
    ('fig_operator_latency.png',
     'GPU-native FEM solver vs. trained-operator latency across batch sizes, '
     'hardware-matched (Tables 10/10a/10b).',
     ['Table 10', 'Table 10a', 'Table 10b']),
    ('fig_breakeven.png',
     "Break-even, in new problem instances, against CPU-FEM and GPU-native "
     "FEM baselines (Table 10d, reusing Table 10c's own GPU-FEM columns).",
     ['Table 10c', 'Table 10d']),
    ('fig_ood_degradation.png',
     'In-distribution vs. out-of-distribution validation error and '
     'degradation factor, all six cases (Table 11).',
     ['Table 11']),
    ('fig_b2_fix_history.png',
     'B2 fix-attempt history (Neo-Hookean) and final adopted B2 results '
     'across materials (Tables 13/14).',
     ['Table 13', 'Table 14']),
    ('fig_physical_quantities.png',
     'Trained-operator displacement, stress, and reaction-force error, all '
     'six cases (Tables 15/16/17).',
     ['Table 15', 'Table 16', 'Table 17']),
    ('fig_operator_vs_fem.png',
     'Operator vs. FEM accuracy and speed-up across mesh resolutions, all '
     'six cases (Tables 18/18a-e).',
     ['Table 18', 'Table 18a', 'Table 18b', 'Table 18c', 'Table 18d', 'Table 18e']),
    ('fig_table20_scaling.png',
     'GPU-native matrix-free solver scaling: wall-clock time and CG '
     'iteration count vs. problem size (Table 20/20a/20c).',
     ['Table 20', 'Table 20a', 'Table 20c']),
    ('fig_torchfem_comparison.png',
     'GPU-native solver vs. torch-fem: solve time and torch-fem peak memory '
     'across resolutions (Table 20d).',
     ['Table 20d']),
    ('fig_pi_vs_dd.png',
     'Physics-informed vs. data-driven training: validation error by loss '
     'and total cost before first inference (Tables 21/21a).',
     ['Table 21', 'Table 21a']),
    ('fig_mms_convergence.png',
     'Method of manufactured solutions: convergence rates, all three '
     'materials, Q4 and Q9 (Tables 22/22a/22b/23/23a).',
     ['Table 22', 'Table 22a', 'Table 22b', 'Table 23', 'Table 23a']),
    ('fig_zeroshot_resolution.png',
     'Zero-shot resolution invariance: mean relative L2 error vs. test '
     'resolution, B1 and B2, all three materials (Table 12).',
     ['Table 12 (revised)', 'Table 12']),
]

# Torch-fem has no dedicated table in the Summary (it's prose-only there);
# resolved by searching for its own discussion text instead of a caption.
SUMMARY_TEXT_ANCHORS = {
    'fig_torchfem_comparison.png': 'torch-fem wins wall-clock by a large',
}


def resolve_all_anchors(doc, tmap, items, is_summary):
    resolved = []  # (fname, base_caption, anchor_elem, anchor_idx)
    for fname, base_caption, labels in FIGURES:
        if is_summary and fname in SUMMARY_TEXT_ANCHORS:
            anchor_elem, idx = text_anchor(items, SUMMARY_TEXT_ANCHORS[fname])
        else:
            anchor_elem, idx = group_anchor(items, tmap, labels)
        resolved.append([fname, base_caption, anchor_elem, idx])
    resolved.sort(key=lambda r: r[3])
    return resolved


def process(docx_in, docx_out, is_summary):
    doc, tmap = build_table_map(docx_in)
    items = load_items(doc)
    resolved = resolve_all_anchors(doc, tmap, items, is_summary)

    for i, (fname, base_caption, anchor_elem, idx) in enumerate(resolved):
        fig_num = START_FIGURE_NUMBER + i
        caption_text = f'Figure {fig_num}. {base_caption}'
        img_path = os.path.join(FIG, fname)
        assert os.path.isfile(img_path), img_path
        insert_figure(anchor_elem, doc, img_path, caption_text)
        print(f'  Figure {fig_num}: {fname}  (anchor idx {idx})')

    doc.save(docx_out)
    print(f'Saved {docx_out} ({len(resolved)} figures inserted)')


if __name__ == '__main__':
    print('=== Report ===')
    process(REPORT, os.path.join(DELIV, 'PFEM_Transolver_Report_updated_2026-09-09.docx'),
            is_summary=False)
    print('\n=== Summary ===')
    process(SUMMARY, os.path.join(DELIV, 'PFEM_Work_Summary_updated_2026-09-09.docx'),
            is_summary=True)
