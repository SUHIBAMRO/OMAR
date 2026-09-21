"""Round 14 -- Timon item 4 (design + build a harder, more realistic 3D
example): first full status report on both candidates prepared in
parallel (B3, a rocking rubber-mount bushing; and a tire sector, a
deliberately lightweight preliminary alternative), including Omar's own
detailed 11-point technical review of both before this write-up, the
real GPU mesh-convergence result that followed, and an honest account of
the field-error methodology fix.

REVISED 2026-09-21b -- Omar's own line-by-line review of the first
draft of this section, before it goes anywhere near Timon, caught seven
real issues, five of them called necessary:
  1. 11.4 did not say explicitly what remains AFTER the candidate is
     picked (dataset generation, training, and the FEM-vs-operator
     comparison itself) -- added.
  2/3. The required-resolution table's own "1%" cell for the region-
     Cauchy average conflated two different things (reference-to-
     reference convergence vs. an operational mesh actually reaching
     1%) and the prose called 0.506% "under" a 0.5-0.7% band when it is
     inside it -- both corrected with Omar's own suggested wording.
  4. THE important one: the ~17.3% "plateau" in 11.2.1 was written up
     as a finding before the one symmetric comparison Omar had asked
     for (element-averaged on BOTH sides, not just one) had actually
     been tried. It has now been implemented in mesh_convergence_B3.py
     and tested directly on real data: it converges cleanly (48.5% ->
     23.3% -> 14.3% -> 7.7% for 144/600/1,568/3,240-element cases
     against the same 9,464-element reference), confirming the earlier
     asymmetric comparison -- not a genuine physical limit -- produced
     the plateau.
  5. Overclaimed language ("mostly discretization noise, not physical
     signal") and an unqualified single-cause claim for the region-
     shrinking test are both softened to what the evidence actually
     supports.
  6. The groove's radius of curvature (0.0912) had no stated units --
     now explicit that this is in the same nondimensional length units
     used throughout the benchmark.
  7. The tire's inflation-pressure surface is now stated explicitly
     (the entire outer boundary of the meridian cross-section, not an
     ambiguous "tread"), and a sentence is added noting the sector-cut
     treatment will be revisited if this candidate is ever selected.

REVISED 2026-09-21d -- the GPU re-run point 4 asked for is now IN: with
the corrected, symmetric methodology, the full Cauchy-tensor field
error between the two finest references (243,360 vs. 424,128 elements)
is 0.662% -- confirming the fix, not just at CPU scale. The whole
resolution ladder now converges cleanly and monotonically (28.442% at
600 elements down to 1.510% at 123,008 elements, 0.662% between the two
references), essentially the same well-behaved shape as the region-
average statistic. 11.2's table and 11.2.1's own text are updated with
these real numbers; nothing about this QoI is provisional any more.

REVISED 2026-09-21e -- Omar's final pass before this goes to the
advisor, four more corrections:
  1. "Region-Cauchy average" and "99th percentile" never stated WHICH
     stress quantity they are -- confirmed in mesh_convergence_B3.py
     (`sigma_flat[:, 0, 0]`) that both are specifically the sigma_xx
     component, not a tensor norm or von Mises stress. Every mention
     now says "sigma_xx" explicitly, and the table's own row labels are
     renamed accordingly. The separate full-tensor field error (all 9
     components) is now labeled distinctly so it is not confused with
     the sigma_xx-only statistics.
  2. "Reaction force" never stated what is actually compared -- confirmed
     in the code (`np.linalg.norm(r["reaction_force"])`) that the
     relative error is computed on the RESULTANT (Euclidean norm) of the
     reaction-force vector, not a single component. Stated explicitly
     everywhere this QoI is named.
  3. Language written for an internal log, not an advisor, is replaced
     throughout with direct scientific phrasing -- e.g. "developed to
     full rigor" -> "developed to a validated FEM mesh-convergence
     stage"; the 11.1 heading "...a real, user-caught geometry
     correction" -> "Geometry refinement and design history"; "honestly
     characterized" -> "validated and documented"; "the project's own
     standing discipline" removed in favor of stating the fact plainly;
     references to whose review caught what are removed from the
     advisor-facing text (they belong in the internal project record,
     not this document).
  4. The tire's pressure load was called an "internal inflation
     pressure" representing the "inner surface of the tire's own outer
     shell" -- but the meridian cross-section here is confirmed
     (`Rr = R_bead + (R_tread_eff - R_bead) * t`, t in [0,1]) to be a
     SOLID volume from the bead to the tread surface, not a hollow shell
     with an internal cavity. Calling this an "inflation pressure" is
     not geometrically accurate. Renamed throughout to "a distributed
     outward pressure preload," with an explicit note that the geometry
     is solid, not a shell, so this is a boundary condition rather than
     a literal internal inflation.

Appended as a new top-level section (11.) after the existing Section 10
(Conclusion and Next Steps) -- this work started after that conclusion
was written and is its own, still-in-progress deliverable, not a
retroactive edit to the existing six-case study.
"""
import os

from docx import Document

DELIV = '/home/user/OMAR/advisor_feedback'
REPORT_SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-19b.docx')
REPORT_DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-21e.docx')

doc = Document(REPORT_SRC)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def para(text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    return p


def table(header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = 'Normal Table'
    for j, htext in enumerate(header):
        cell = tbl.rows[0].cells[j]
        cell.text = htext
        cell.paragraphs[0].runs[0].bold = True
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            tbl.rows[i + 1].cells[j].text = str(val)
    return tbl


# =======================================================================
h1("11. Third Benchmark Candidate (in progress): 3D Hyperelastic "
   "Rubber-Mount Bushing vs. Tire Sector")

para(
    "Timon's item 4 asked for a new, harder, more realistic 3D example beyond "
    "the B1/B2 benchmark pair. Two candidates were prepared in parallel, to "
    "the same single-material (Neo-Hookean), no-contact scope, so Timon can "
    "choose between them before any expensive commitment: (A) a rocking "
    "elastomeric rubber-mount bushing (\"B3\"), developed to a validated FEM "
    "mesh-convergence stage with a GPU mesh-convergence study; and (B) a tire "
    "sector -- a torus segment, not a straight extrusion -- prepared as a "
    "deliberately lightweight preliminary alternative (geometry, boundary "
    "conditions, a solve, and a basic convergence check only, not yet brought "
    "to the same validated mesh-convergence stage as B3). Dataset generation "
    "and neural-operator training have not started for either candidate; "
    "that step depends on this choice."
)

h2("11.1 Geometry refinement and design history")

para(
    "B3's geometry went through two substantive revisions before reaching "
    "its current form, both identified through detailed review of the "
    "design. The first draft was a "
    "thin plate with a circular hole under uniaxial tension -- rejected as "
    "both the wrong shape category (not the \"rubber mount\" character "
    "requested) and too simple a benchmark. The corrected design is an "
    "elastomeric bushing: a hollow rubber annulus bonded to a rigid outer "
    "housing (fixed) and a rigid inner core given a prescribed rocking "
    "rotation. The second correction addressed the source of the stress "
    "concentration itself: an initial fillet-terminated bond still carried "
    "a sharp bonded-to-free boundary-condition discontinuity at the edge of "
    "the fillet band, which can itself reproduce a free-edge stress "
    "singularity independent of how smooth the surrounding geometry looks. "
    "The final design bonds the core to the rubber CONTINUOUSLY over its "
    "full height, with no free/bonded transition anywhere on the inner "
    "surface; the one explicit, disclosed, finite-radius feature is a "
    "smooth (C1-continuous, raised-cosine) circumferential GROOVE in the "
    "core's own radius profile at mid-height, with a closed-form radius of "
    "curvature of 0.0912 in the same nondimensional length units used "
    "throughout this benchmark (R_in0=0.5, R_out=1.0, Lz=1.0 -- consistent "
    "with this project's own dimensionless geometry/material convention "
    "already used for B1/B2, not tied to a specific physical unit). The "
    "core's prescribed motion is a true rigid-body rotation (coupled u_x "
    "and u_z), not a linear small-angle approximation, which was verified "
    "to silently assume zero axial displacement -- wrong for a real tilt."
)

h2("11.2 Candidate A: the B3 rubber-mount bushing -- GPU-confirmed "
   "mesh convergence")

para(
    "Full QoI set tracked at every resolution: displacement L2 and an "
    "H1-like gradient semi-norm (both against a fine reference, via "
    "interpolation in the mesh's own shared parametric coordinate "
    "space, exact for node fields); total hyperelastic strain energy "
    "(a scalar total -- explicitly NOT the same, more rigorous "
    "tangent-energy-norm metric B1/B2 use, see the correction below); "
    "reaction moment about the rotation axis (the PRIMARY reaction QoI for "
    "this rocking case) and the resultant (Euclidean norm) of the reaction "
    "force vector (secondary, confirmed non-degenerate from real data, not "
    "assumed); and a fixed-region statistic on the sigma_xx component of the "
    "Cauchy stress tensor at the groove -- a volume-weighted regional "
    "average of sigma_xx and a regional 99th percentile of sigma_xx -- plus "
    "a separate full Cauchy-TENSOR relative field error over all stress "
    "components (see 11.2.1 for its own corrected methodology), with a "
    "reliability gate that reports the sigma_xx 99th percentile as \"not "
    "reliable\" below 20 quadrature-point samples rather than a misleading "
    "number from too few points."
)

para(
    "A GPU run (A100) solved two fine references to check whether the "
    "hardest QoI -- the fixed-region sigma_xx statistic -- had actually "
    "converged, not merely reached the largest mesh tried: 243,360 "
    "elements (81 x 40 x 79) and a new, finer 424,128-element mesh "
    "(97 x 48 x 95). The volume-weighted regional average of sigma_xx "
    "changed by only 0.506% between these two references -- down from "
    "0.87% at the previous step, and within the predefined approximately "
    "0.5-0.7% reference-convergence band -- supporting that this "
    "statistic's reference value is itself converged. This is a distinct "
    "claim from any operational mesh reaching 1% error against that "
    "reference: the highest non-reference mesh in the tested ladder "
    "(123,008 elements) still shows 1.28% error, so 1% is not yet reached "
    "by any mesh actually tested, only supported as a reasonable target "
    "by the reference's own demonstrated stability. The required-"
    "resolution table below is computed against the newer, finer "
    "reference."
)

table(
    ["QoI", "5% threshold", "2% threshold", "1% threshold"],
    [
        ["Displacement L2", "600 el.", "3,240 el.", "20,808 el."],
        ["H1 (gradient) semi-norm", "38,808 el.", "not reached", "not reached"],
        ["Total strain energy", "600 el.", "600 el.", "3,240 el."],
        ["Reaction moment (primary)", "3,240 el.", "9,464 el.", "38,808 el."],
        ["Reaction force resultant (Euclidean norm, secondary)", "65,000 el.", "not reached", "not reached"],
        ["Regional average of sigma_xx (volume-weighted)", "38,808 el.", "123,008 el.",
         "not reached by any non-reference mesh in the tested ladder "
         "(123,008 elements gives 1.28%); the two fine references differ "
         "by only 0.506%, supporting convergence of the reference itself"],
        ["Regional 99th percentile of sigma_xx (volume-weighted)", "not reached", "not reached", "not reached"],
        ["Cauchy-stress tensor field error (all components)", "38,808 el.", "123,008 el.",
         "not reached by any non-reference mesh in the tested ladder "
         "(123,008 elements gives 1.51%); the two fine references differ "
         "by only 0.662%, supporting convergence of the reference itself "
         "(see 11.2.1 for the methodology correction that produced this "
         "cleanly converging result)"],
    ],
)

para(
    "Displacement, energy, and reaction moment all converge cleanly and "
    "smoothly, which is reassuring evidence the model, mesh, and boundary "
    "conditions are correct rather than merely not visibly broken. The H1 "
    "semi-norm and the reaction-force resultant do not reach 1% within the "
    "tested range -- a disclosed limitation, consistent with this project's "
    "own repeated finding (already established for B1/B2) that gradient/"
    "energy-flux-type quantities converge more slowly than displacement or "
    "averaged-stress statistics.",
    italic=True,
)

h2("11.2.1 The Cauchy-tensor field-error metric: a methodology bug "
   "found, fixed, and confirmed on real GPU data")

para(
    "An earlier draft of this section reported the full Cauchy-TENSOR "
    "relative field error (comparing the whole stress tensor pointwise, at "
    "every quadrature point in the fixed region, rather than a single "
    "averaged scalar) as plateauing around 17.3% between the two finest GPU "
    "references, and treated that as a converged finding. On review, this "
    "was premature: the comparison itself was ASYMMETRIC -- the coarser "
    "mesh's own per-ELEMENT-AVERAGED Cauchy field was compared against the "
    "reference's own RAW per-Gauss-point values, two different "
    "representations of the same field, not a fair apples-to-apples test."
)

para(
    "Two fixes were tried before the correct one was found, all tested "
    "directly against the real solver rather than assumed to work. (1) "
    "Moving BOTH sides to raw per-Gauss-point values made the reported "
    "error substantially WORSE (a 600-element case moved from 32.97% to "
    "84.08%; a 3,240-element case from 22.80% to 59.51%), traced to a real, "
    "checkable cause: within a single coarse element, the stress at its own "
    "8 Gauss points varies by up to 250.9 -- comparable to or larger than "
    "the entire mesh's own element-averaged stress range end to end. Raw "
    "Gauss-point stresses are generally discontinuous across element "
    "boundaries and can be highly mesh-sensitive in regions with steep "
    "stress gradients, so comparing raw values on both sides amplified "
    "that sensitivity rather than removing it. (2) Shrinking the fixed "
    "sampling region (to test whether a smaller zone would avoid spanning "
    "as steep a gradient) was tried at half and quarter the original "
    "radius: the reported error increased at every step tested (48.78% to "
    "58.04% to 59.73%) rather than improving. A smaller region also means "
    "fewer quadrature-point samples available to average over, which may "
    "contribute to this, though that was not isolated as the single "
    "confirmed cause. Neither (1) nor (2) was adopted."
)

para(
    "The combination not yet tried at that point was the genuinely "
    "SYMMETRIC one: both sides at the SAME (element-averaged) "
    "representation, compared at the SAME physical sample points -- the "
    "coarser case's own field interpolated onto the reference's own "
    "ELEMENT CENTROIDS (not Gauss points), compared against the "
    "reference's own element-averaged field at those same centroids, "
    "volume-weighted. This has now been implemented in "
    "mesh_convergence_B3.py and tested directly on real CPU data: it shows "
    "clean, monotonic convergence with resolution -- 48.5%, 23.3%, 14.3%, "
    "and 7.7% for the 144-, 600-, 1,568-, and 3,240-element cases, all "
    "against the same 9,464-element reference. This is a real, physically "
    "sensible convergence trend, unlike either asymmetric version tried "
    "before it, and supports that the earlier ~17.3% plateau was a "
    "methodology artifact rather than a genuine physical limit of this QoI."
)

para(
    "This fix has now been confirmed on real GPU data, not just at CPU "
    "scale. Re-running the same A100 reference pair (243,360 vs. 424,128 "
    "elements) with the corrected, symmetric comparison gives a full "
    "Cauchy-tensor field error of 0.662% between the two references -- "
    "and the WHOLE resolution ladder now converges cleanly and "
    "monotonically against the finer reference: 28.442% (600 elements), "
    "15.130% (3,240), 8.609% (9,464), 5.567% (20,808), 3.762% (38,808), "
    "2.599% (65,000), 1.510% (123,008), down to 0.662% between the two "
    "references -- essentially the same well-behaved shape as the "
    "regional sigma_xx average (11.2 above), and a world away from the "
    "earlier ~17.3% plateau. This QoI is no longer provisional: the "
    "required-resolution table above now reports real 5%/2% thresholds "
    "for it, with 1% not yet reached by any single non-reference mesh but "
    "well-supported by the reference's own 0.662% stability, exactly "
    "mirroring how the regional sigma_xx average's own 1% row is read.",
    italic=True,
)

h2("11.3 Candidate B: the tire sector -- a lightweight preliminary "
   "alternative")

para(
    "A genuine torus segment (not a straight extrusion, which would only "
    "reproduce B3): the same meridian half-ring cross-section B2/B3 "
    "already use is swept through a limited circumferential sector angle "
    "around a big wheel axis. The bead is bonded to a fixed rigid rim; the "
    "tread carries a smooth (C1, raised-cosine) groove at its own "
    "centerline, the same construction as B3's own feature. Two real bugs "
    "were found and fixed during development: a parity-flipping axis "
    "permutation in the torus-sweep mapping gave every element negative "
    "signed volume (fixed by reversing the underlying 2D quad's own node "
    "winding); and an inversion of the sweep mapping in the boundary-mask "
    "logic swapped which recovered local coordinate was which, making the "
    "tread-load selection always come up empty (fixed by re-deriving the "
    "inverse mapping directly from the forward one)."
)

para(
    "Loading combines two superposed pressure loads under one incremental "
    "ramp: a distributed outward pressure preload, applied over the entire "
    "outer boundary of the meridian cross-section (every node with "
    "r_local = R_tread_eff(theta) across the full theta range, i.e. every "
    "node not bonded to the bead). The meridian cross-section here is a "
    "SOLID volume from the bead to the tread surface, not a hollow shell "
    "with an internal cavity, so this preload is a boundary condition on "
    "the outer surface rather than a literal internal inflation pressure -- "
    "the term \"inflation pressure\" used in an earlier draft has been "
    "corrected accordingly. A second, additional, localized inward pressure "
    "is superposed within a window at the tread centerline (renamed from an "
    "earlier, inaccurate \"contact patch\" -- there is no ground-contact "
    "formulation here, no contact mechanics, no rigid ground surface). A "
    "solve smoke test confirms the rim stays exactly fixed, the rest of the "
    "tread bulges outward under the distributed preload alone as expected, "
    "the loaded window still moves net inward (the local load dominates the "
    "preload there), and the deformation is non-degenerate in all three "
    "directions, confirming the torus topology is genuinely three-"
    "dimensional."
)

para(
    "A lightweight preliminary convergence check (four resolutions, 96 to "
    "2,880 elements) tracks displacement, total strain energy, and the "
    "groove-region stress; both show real, material resolution "
    "sensitivity, consistent with this project's own established pattern "
    "that local, boundary-window-dependent quantities converge more slowly "
    "than global ones. Two disclosed effects were identified and handled "
    "explicitly rather than reported without qualification: the reaction "
    "moment about the wheel's own spin axis is exactly zero at every "
    "resolution, for a genuine geometric reason (every applied load is "
    "pressure normal to a surface of revolution about that axis, which "
    "cannot produce torque about it) -- reported as a diagnostic only, "
    "never as a convergence target, unlike B3's own physical rocking "
    "moment. The combined net reaction force (the distributed preload and "
    "local load together) swings by roughly 100x with sign changes across "
    "resolutions because it is a near-cancellation of two comparable, "
    "unrelated load resultants -- confirmed by reconstructing the combined "
    "force exactly from the two loads' own isolated resultants -- so the "
    "local load's own isolated resultant magnitude is used as the reaction-"
    "family QoI instead, and the combined force is kept only for the "
    "equilibrium check. The two circumferential sector-cut faces are "
    "confirmed genuinely free (not artificially clamped), but are not tied "
    "together as a periodic boundary condition -- a disclosed limitation, "
    "since implementing a true periodic-tie constraint is nontrivial new "
    "solver infrastructure judged out of scope for a lightweight "
    "preliminary candidate; the load and stress-sampling region are "
    "deliberately centered mid-sector, away from both cuts, to mitigate "
    "this. If the tire candidate is selected for the final benchmark, the "
    "sector-boundary treatment will be revisited before dataset generation "
    "or operator training -- this preliminary version is not proposed as "
    "the final tire model."
)

h2("11.4 Status")

para(
    "Both candidates are prepared to a validated, documented state -- B3 "
    "with a GPU-confirmed mesh-convergence study and a full required-"
    "resolution table; the tire sector as a working, but deliberately "
    "lightweight preliminary alternative. No dataset generation or "
    "neural-operator training has started for either candidate; that step "
    "depends on Timon's choice between them. After the final 3D candidate "
    "is selected, the remaining numerical step is dataset generation, "
    "operator training, and the FEM-versus-operator comparison on "
    "displacement, reaction, energy, and regional Cauchy-stress QoIs, "
    "exactly as done for B1 and B2 -- that comparison has not started for "
    "either candidate and is not implied to be complete by anything above."
)

doc.save(REPORT_DST)
print('Saved', REPORT_DST)

check = Document(REPORT_DST)
print('paragraphs:', len(check.paragraphs), '(was 626)')
print('tables:', len(check.tables), '(was 102)')
print('images:', len(check.inline_shapes))
