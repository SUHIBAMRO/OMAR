"""Companion to add_qoi_crossover_remaining_cases_to_report.py -- same
round-11 point 2 answer, condensed for the Summary: which QoI binds the
coarsest-suitable-FEM crossover for the five remaining cases."""
import copy
import os

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-17d.docx')
DST = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-17e.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_table_after(anchor_para, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    pr = doc.tables[0]._tbl.find(qn('w:tblPr'))
    if pr is not None:
        old = tbl._tbl.find(qn('w:tblPr'))
        if old is not None:
            tbl._tbl.remove(old)
        tbl._tbl.insert(0, copy.deepcopy(pr))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    tbl._tbl.getparent().remove(tbl._tbl)
    anchor_para._p.addnext(tbl._tbl)
    return tbl


def insert_after_table(table, style_para, text):
    new_p_el = copy.deepcopy(style_para._p)
    table._tbl.addnext(new_p_el)
    np = Paragraph(new_p_el, style_para._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


anchor = find_para_exact(
    'Figure R10-F. Peak-stress QoI vs. displacement error, five cases (Table R10-6).'
)

p1 = insert_after(
    anchor,
    "Round-11 point 2: does the same QoI determine the coarsest-suitable-FEM "
    "crossover every time? Already answered for B1xNeo-Hookean (N=11, bound by "
    "tangent energy) -- repeated here for the other five cases (Table R10-7). "
    "Checkpoint caveat as above: B1xMooney-Rivlin, B1xArruda-Boyce, "
    "B2xMooney-Rivlin and B2xNeo-Hookean reflect their ORIGINAL pre-retrain "
    "checkpoints here (task #22 predates all four retrains)."
)

HEADER = ['Case', 'Coarsest suitable N', 'Binding QoI']
t = insert_table_after(p1, HEADER, [
    ['B1 x Mooney-Rivlin', '3', 'all five metrics at once (L2, H1, energy, peak stress, reaction)'],
    ['B1 x Arruda-Boyce', '3', 'all five metrics at once (L2, H1, energy, peak stress, reaction)'],
    ['B2 x Neo-Hookean', '4', 'tangent energy'],
    ['B2 x Mooney-Rivlin', '4', 'tangent energy'],
    ['B2 x Arruda-Boyce', '4', 'tangent energy'],
])

caption_para = insert_after_table(
    t, p1,
    'Table R10-7. Coarsest suitable FEM and its binding QoI, five cases, N=1401. '
    'N=3 is the coarsest resolution tested, so the two B1 rows\' true crossover '
    'may lie even coarser.'
)

insert_after(
    caption_para,
    "Mixed answer: all three B2 materials share the SAME binding metric (tangent "
    "energy), matching B1xNeo-Hookean -- but B1's other two materials cross on "
    "EVERY metric simultaneously, at the coarsest mesh tested (N=3), because the "
    "pre-retrain operator's own accuracy there is poor enough (L2_rel 33.65% "
    "Mooney-Rivlin, 44.57% Arruda-Boyce) that even the cheapest FEM mesh in this "
    "report already wins on every quantity of interest."
)

doc.save(DST)
print('Saved', DST)
