"""Adds the peak-stress QoI numbers (Task #22, "extend N=1401 accuracy/
QoI/break-even analysis to all 6 cases") for the five cases beyond
B1xNeo-Hookean -- computed with real GPU numbers on 2026-09-15 but
never transcribed into the Report itself, per the 2026-09-16 audit
(Gap 2).

IMPORTANT CHECKPOINT-VERSION CAVEAT, stated explicitly in the added
text rather than left implicit: these five numbers were measured on
2026-09-15, BEFORE this week's own multi-resolution retrains for
B1xMooney-Rivlin, B1xArruda-Boyce and B2xMooney-Rivlin. Their
disp_rel_L2 values here (39.20%, 45.62%, 49.47%) match exactly the
"OLD" checkpoint column already published in Tables 18-R10f/g/h, which
confirms these three rows reflect the ORIGINAL (pre-retrain)
checkpoints, not the newly retrained ones -- the peak-stress metric
itself has not been recomputed against the new checkpoints for any of
the three. B2xNeo-Hookean (still mid-training its own first-ever
multi-res retrain) and B2xArruda-Boyce (no retrain planned, excluded
per Omar's own explicit choice) have only ever had one checkpoint each,
so no such caveat applies to those two rows.

Inserted directly after the Arruda-Boyce root-cause paragraph (the tail
of the resolution-matched break-even addition) and before "The same
comparison for the other two B1 materials..." -- keeping every N=1401,
all-six-cases result from task #21/#22 together in one place.
"""
import copy
import os

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-16c.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-16d.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_table_after(anchor_para, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    pr = doc.tables[0]._tbl.find(qn('w:tblPr'))
    if pr is not None:
        old = tbl._tbl.find(qn('w:tblPr'))
        if old is not None:
            tbl._tbl.remove(old)
        tbl._tbl.insert(0, copy.deepcopy(pr))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    tbl._tbl.getparent().remove(tbl._tbl)
    anchor_para._p.addnext(tbl._tbl)
    return tbl


def insert_after_table(table, style_para, text):
    new_p_el = copy.deepcopy(style_para._p)
    table._tbl.addnext(new_p_el)
    np = Paragraph(new_p_el, style_para._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


anchor = find_para_exact(
    "Both Arruda-Boyce cases fail at N=1401 in every run attempted so far, for a "
    "genuine and already root-caused reason distinct from the memory-cleanup issue "
    "above: torch-fem's own Newton-Raphson solve does not converge (\"did not "
    "converge in increment 8 after 10 cutbacks\"), and the real underlying cause, "
    "read directly from the exception's own cause chain rather than assumed, is a "
    "CUDA out-of-memory error inside torch-fem's own Hessian assembly for this "
    "material specifically -- Arruda-Boyce's own strain-energy density is more "
    "expensive to differentiate twice than Neo-Hookean's or Mooney-Rivlin's, and at "
    "N=1401's problem size that cost alone exhausts the GPU before torch-fem's own "
    "solve can complete. This is reported as a real finding about torch-fem's own "
    "scaling limit for this material at this resolution, not a defect in this "
    "project's own comparison code."
)

p1 = insert_after(anchor,
    "Task #22 also asked for the peak-stress QoI (fixed physical location, "
    "converged from a fine reference -- the same convention as B1xNeo-Hookean's "
    "own peak-stress discussion above) at N=1401 for the remaining five cases. "
    "Table 18-R10k. A checkpoint-version caveat applies to three of the five rows: "
    "B1xMooney-Rivlin, B1xArruda-Boyce and B2xMooney-Rivlin were measured here "
    "against their ORIGINAL (pre-retrain) checkpoints -- their disp_rel_L2 values "
    "match exactly the \"OLD\" column already published in Tables 18-R10f/g/h -- so "
    "the peak-stress numbers below have not been recomputed against this week's "
    "own newly retrained checkpoints for those three cases. B2xNeo-Hookean (its own "
    "first multi-res retrain still mid-training) and B2xArruda-Boyce (no retrain "
    "planned) have only ever had one checkpoint each, so no such caveat applies to "
    "those two rows.")

HEADER = ['Case', 'disp_rel_L2 @N=1401', 'peak_stress_rel_err @N=1401']
t = insert_table_after(p1, HEADER, [
    ['B1 x Mooney-Rivlin', '39.20%', '78.05% (original checkpoint)'],
    ['B1 x Arruda-Boyce', '45.62%', '78.25% (original checkpoint)'],
    ['B2 x Neo-Hookean', '46.37%', '49.48%'],
    ['B2 x Mooney-Rivlin', '49.47%', '48.13% (original checkpoint)'],
    ['B2 x Arruda-Boyce', '39.73%', '39.97%'],
])
p2 = insert_after_table(t, p1,
    "Table 18-R10k. Peak-stress QoI at N=1401, the five cases beyond "
    "B1xNeo-Hookean (already covered above). Same fixed-location, fine-reference "
    "convention as the B1xNeo-Hookean peak-stress discussion.")

p3 = insert_after(p2,
    "Two patterns are worth stating plainly rather than left for the reader to "
    "notice. First, B1's own two cases here (78.05%, 78.25%) are structurally "
    "worse on this metric than every B2 case (39.97%-49.48%) -- consistent with "
    "the B1xNeo-Hookean finding above that the located peak-stress point sits near "
    "a boundary-condition-transition corner prone to a stress singularity, a "
    "property of the B1 geometry's own boundary conditions rather than of any "
    "particular material. Second, B2xArruda-Boyce (39.97%) has the BEST "
    "peak-stress accuracy of all six cases in this report, despite being one of "
    "the two cases that fails outright on the resolution-matched break-even above "
    "-- the two findings are unrelated: the break-even failure is torch-fem's own "
    "reference solve running out of memory, not a property of the operator's own "
    "accuracy, which this table shows is genuinely strong for this case.")

doc.save(DST)
print('Saved', DST)
