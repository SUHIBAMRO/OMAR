"""Fixes a real textual inconsistency Omar caught (2026-09-18): the
round-12 point 2 intro paragraph said "six multi-resolution retrains
plus the direct-N1401 ablation" (seven), but the very same table's own
caption right after it correctly says B2xArruda-Boyce has NO
multi-resolution retrain at all -- only five cases were actually
retrained on the multi-resolution recipe, plus B2xArruda-Boyce's own
original run, plus the direct-N1401 ablation, is what actually makes
seven. Applies the same fix to all three documents that carry this
paragraph verbatim (Report, round-12-only Summary, side document).
"""
import os

from docx import Document

DELIV = '/home/user/OMAR/advisor_feedback'
TARGETS = [
    os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-18.docx'),
    os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-18b.docx'),
    os.path.join(DELIV, 'Round12_Reply_to_Timon_Points_2026-09-18.docx'),
]

OLD = (
    "Table 18-R10n reports exactly that for all seven completed training runs (six "
    "multi-resolution retrains plus the direct-N1401 ablation just discussed), read "
    "directly from each run's own already-saved metrics_history.json -- nothing "
    "retrained for this table."
)
NEW = (
    "Table 18-R10n reports exactly that for all seven completed training runs: five "
    "multi-resolution retrains, one original B2 x Arruda-Boyce run (it has no "
    "multi-resolution retrain of its own -- see the table's own caption), and the "
    "direct-N1401 ablation just discussed -- read directly from each run's own "
    "already-saved metrics_history.json, nothing retrained for this table."
)


def replace_paragraph_text(para, new_text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(new_text)


for path in TARGETS:
    doc = Document(path)
    hits = [p for p in doc.paragraphs if OLD in p.text]
    assert len(hits) == 1, f'{path}: expected exactly 1 match, found {len(hits)}'
    full_text = hits[0].text
    assert full_text.count(OLD) == 1
    replace_paragraph_text(hits[0], full_text.replace(OLD, NEW))
    doc.save(path)
    print('Fixed', path)
