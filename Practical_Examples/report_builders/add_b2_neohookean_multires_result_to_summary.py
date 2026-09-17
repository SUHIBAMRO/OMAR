"""Companion to add_b2_neohookean_multires_result_to_report.py -- same
real result (B2 x Neo-Hookean's multi-res retrain finished, 46.37% ->
36.05% at N=1401), added to the Summary's condensed prose, its own
cross-case table (Table R10-2), and its own figure (Figure R10-B)."""
import copy
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

FIG = '/home/user/OMAR/Practical_Examples/report_builders/figures'
DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-17c.docx')
DST = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-17d.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def replace_paragraph_text(para, new_text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(new_text)
    return para


def insert_figure_after(anchor_para, image_path, caption_text, width_in=5.5):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width_in))
    img_elem = img_p._p
    img_elem.getparent().remove(img_elem)

    cap_p = doc.add_paragraph()
    cap_run = cap_p.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_elem = cap_p._p
    cap_elem.getparent().remove(cap_elem)

    anchor_para._p.addnext(cap_elem)
    anchor_para._p.addnext(img_elem)
    return Paragraph(cap_elem, anchor_para._parent)


def add_table_row(table, values):
    new_tr = copy.deepcopy(table.rows[-1]._tr)
    table._tbl.append(new_tr)
    new_row = table.rows[-1]
    for j, v in enumerate(values):
        cell = new_row.cells[j]
        for p in cell.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
        cell.paragraphs[0].add_run(str(values[j]))
    return new_row


# 1) Prose: replace the "still mid-training" sentence with the real result.
prose = find_para_exact(
    "Results, N=1401, old vs. new checkpoint (Table R10-2): B1 x Mooney-Rivlin "
    "39.20% -> 15.04% (~62% relative reduction); B1 x Arruda-Boyce 45.62% -> 25.05% "
    "(~45%, and its own real GPU run is reassuring evidence that a shared "
    "chain-locking-clamp exposure with this project's own torch-fem Arruda-Boyce "
    "runs does not bite in practice here -- all 16 resolutions, both checkpoints, "
    "converged cleanly); B2 x Mooney-Rivlin 49.47% -> 13.10% (~73.5%, the largest "
    "reduction so far, and the clearest case of a genuinely FLAT error across "
    "N=37-1401 rather than merely a smaller one at N=1401). B2 x Neo-Hookean's own "
    "retrain was still mid-training as of this writing and is not included below."
)
replace_paragraph_text(
    prose,
    "Results, N=1401, old vs. new checkpoint (Table R10-2): B1 x Mooney-Rivlin "
    "39.20% -> 15.04% (~62% relative reduction); B1 x Arruda-Boyce 45.62% -> 25.05% "
    "(~45%, and its own real GPU run is reassuring evidence that a shared "
    "chain-locking-clamp exposure with this project's own torch-fem Arruda-Boyce "
    "runs does not bite in practice here -- all 16 resolutions, both checkpoints, "
    "converged cleanly); B2 x Mooney-Rivlin 49.47% -> 13.10% (~73.5%, the largest "
    "reduction so far, and the clearest case of a genuinely FLAT error across "
    "N=37-1401 rather than merely a smaller one at N=1401); B2 x Neo-Hookean 46.37% "
    "-> 36.05% (~22.3%, the smallest of the five reductions -- unlike Mooney-"
    "Rivlin's flat curve, this checkpoint's error is lowest right where it was "
    "trained, N=101 at 12.98%, and rises again beyond that to 36-40% by N=1401, "
    "still clearly better than the original checkpoint's 46.4-46.7% there)."
)

# 2) Table R10-2: add the 5th row.
summary_table = None
for tbl in doc.tables:
    if [c.text for c in tbl.rows[0].cells] == [
            'Case', 'Old (N=1401)', 'New (N=1401)', 'Relative reduction']:
        summary_table = tbl
        break
assert summary_table is not None, 'could not find Table R10-2'
assert len(summary_table.rows) == 5, f'expected 4 data rows before edit, got {len(summary_table.rows) - 1}'
add_table_row(summary_table, ['B2 x Neo-Hookean', '46.37%', '36.05%', '~22.3%'])

# 3) Figure R10-B: replace the embedded image with the 5-case version.
fig_caption = find_para_exact(
    'Figure R10-B. Multi-resolution retrain fix, old vs. new checkpoint, all '
    'completed cases (Table R10-2).'
)
old_img_p = fig_caption._p.getprevious()
assert old_img_p is not None and old_img_p.findall('.//' + qn('w:drawing')), \
    'expected an image paragraph immediately before Figure R10-B\'s caption'
fig_caption_text = fig_caption.text
old_img_p.getparent().remove(old_img_p)
fig_caption._p.getparent().remove(fig_caption._p)

table_caption = find_para_exact('Table R10-2. Multi-resolution retrain fix, every completed case, disp_rel_L2 at N=1401.')
insert_figure_after(table_caption, os.path.join(FIG, 'fig_multires_retrain_summary.png'), fig_caption_text)

doc.save(DST)
print('Saved', DST)
