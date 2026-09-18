"""Rebuilds round-12 point 1 properly: Timon's own wording was "for EACH
resolution, please compare FEM and NO against the same fine reference in
displacement L2, H1/energy norm, reaction force and stress" -- the
previous version of this section (Tables 18-R10o/p/q) did not actually
do that: it gave the full 16-resolution sweep for ONE case (B1xNeo-
Hookean) but ONLY the Cauchy-stress metric there, and gave the other
five cases only a SINGLE-resolution snapshot (N=13), for both metric
groups. Omar caught this (2026-09-18) before sending -- a real,
justified catch, not a nitpick.

This script replaces that whole block with twelve tables: for EACH of
the six cases, one classical-QoI table (L2, H1, energy, reaction, all
sixteen resolutions N=3..49) and one Cauchy-stress table (region average,
99th percentile, true max, same sixteen resolutions) -- genuinely "for
each resolution," for every case, for every metric group Timon asked
for.

Data sources, merged (never hand-transcribed):
  - round12_final_accuracy_cauchy_summary.json (all six cases, FEM+NO,
    L2/H1/energy/reaction/Cauchy-avg/p99/max at all 16 N -- though NO's
    classical L2/H1/energy/reaction is null below N=13 for five of the
    six cases, a genuine, disclosed pre-existing gap: the operator's own
    zero-shot accuracy was simply never evaluated below N=13 before this
    round, for any case).
  - no_accuracy_degradation_sweep_B1_neo_hookean.json (the 2026-09-18
    follow-up run that closed this same gap for the flagship case only,
    giving it full N=3..49 classical-QoI coverage where the other five
    cases still have only N=13..49).

Table numbering: continues the existing 18-R10o/p/q sequence, now one
letter per table across all twelve (18-R10o..18-R10z) -- B1xNeo-Hookean's
own Cauchy table reuses the exact same figure already built for it
(fig_round12_cauchy_flagship.png), since that table's own content did
not change, only its label (from 18-R10o to 18-R10p, to keep the
per-case classical/Cauchy pairing consistent across all six cases).
"""
import copy
import json
import os

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
FIG = '/home/user/OMAR/Practical_Examples/report_builders/figures'
SCRATCH = '/tmp/claude-0/-home-user/64d7c4d8-d5f0-5686-a58f-aa87abfd4ba4/scratchpad/round12'

REPORT = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-18.docx')

# ---------------------------------------------------------------------
# data: merge the B1xNeo-Hookean follow-up sweep into the main summary
# ---------------------------------------------------------------------
with open(os.path.join(SCRATCH, 'round12_final_accuracy_cauchy_summary.json')) as f:
    DATA = json.load(f)
with open(os.path.join(SCRATCH, 'no_accuracy_degradation_sweep_B1_neo_hookean.json')) as f:
    B1NH_FIX = json.load(f)

b1nh_fix_by_n = {r['N']: r['fp32'] for r in B1NH_FIX['rows']}
for row in DATA['cases']['B1_neo_hookean']:
    fx = b1nh_fix_by_n.get(row['N'])
    if fx:
        row['no_l2_rel'] = fx['L2_rel']
        row['no_h1_semi_rel'] = fx['H1_semi_rel']
        row['no_energy_rel'] = fx['energy_rel']
        row['no_reaction_rel_err'] = fx['reaction_resultant_rel_err']

CASE_LABELS = {
    'B1_neo_hookean': 'B1 x Neo-Hookean',
    'B1_mooney_rivlin': 'B1 x Mooney-Rivlin',
    'B1_arruda_boyce': 'B1 x Arruda-Boyce',
    'B2_neo_hookean': 'B2 x Neo-Hookean',
    'B2_mooney_rivlin': 'B2 x Mooney-Rivlin',
    'B2_arruda_boyce': 'B2 x Arruda-Boyce',
}
CASE_ORDER = list(CASE_LABELS)
LOW_N = DATA['low_N']

# table letters 'o'..'z', two per case (classical, cauchy), in order
LETTERS = list('opqrstuvwxyz')
TABLE_ID = {}
for i, case_id in enumerate(CASE_ORDER):
    TABLE_ID[(case_id, 'classical')] = f'18-R10{LETTERS[2 * i]}'
    TABLE_ID[(case_id, 'cauchy')] = f'18-R10{LETTERS[2 * i + 1]}'


def pct(x):
    return f'{x * 100:.2f}%' if x is not None else 'n/a'


CLASSICAL_HEADER = ['N', 'FEM L2', 'Op. L2', 'FEM H1', 'Op. H1', 'FEM Energy', 'Op. Energy', 'FEM Reaction', 'Op. Reaction']
CAUCHY_HEADER = ['N', 'FEM avg', 'Op. avg', 'FEM p99', 'Op. p99', 'FEM max', 'Op. max']


def classical_rows(case_id):
    rows = []
    for r in DATA['cases'][case_id]:
        rows.append([
            str(r['N']),
            pct(r['fem_l2_rel']), pct(r['no_l2_rel']),
            pct(r['fem_h1_semi_rel']), pct(r['no_h1_semi_rel']),
            pct(r['fem_energy_rel']), pct(r['no_energy_rel']),
            pct(r['fem_reaction_rel_err']), pct(r['no_reaction_rel_err']),
        ])
    return rows


def cauchy_rows(case_id):
    rows = []
    for r in DATA['cases'][case_id]:
        rows.append([
            str(r['N']),
            pct(r['fem_cauchy_avg_rel_err']), pct(r['no_cauchy_avg_rel_err']),
            pct(r['fem_cauchy_p99_rel_err']), pct(r['no_cauchy_p99_rel_err']),
            pct(r['fem_cauchy_max_rel_err']), pct(r['no_cauchy_max_rel_err']),
        ])
    return rows


# =======================================================================
# docx surgery helpers (remove a range, then insert fresh content)
# =======================================================================
def body_children(doc):
    return list(doc.element.body)


def elem_text(el):
    if el.tag != qn('w:p'):
        return None
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_index(children, text):
    matches = [i for i, el in enumerate(children)
               if elem_text(el) is not None and elem_text(el).strip() == text]
    assert len(matches) == 1, f'{len(matches)} matches for {text!r}'
    return matches[0]


def _elem_of(x):
    if hasattr(x, '_p'):
        return x._p
    if hasattr(x, '_tbl'):
        return x._tbl
    return x


def remove_range_return_anchor(doc, start_text, end_text):
    """Removes every body element from start_text's paragraph through
    end_text's paragraph (inclusive), and returns the element now
    immediately preceding that gap, to insert new content after."""
    children = body_children(doc)
    i0 = find_index(children, start_text)
    i1 = find_index(children, end_text)
    assert i1 >= i0
    anchor = children[i0 - 1]
    for el in children[i0:i1 + 1]:
        el.getparent().remove(el)
    return anchor


def insert_paragraph_after(doc, anchor, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    p_elem = p._p
    p_elem.getparent().remove(p_elem)
    _elem_of(anchor).addnext(p_elem)
    return p_elem


def insert_table_after(doc, anchor, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    for j, h in enumerate(header):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.rows[i].cells[j].text = v
    tbl_elem = tbl._tbl
    tbl_elem.getparent().remove(tbl_elem)
    _elem_of(anchor).addnext(tbl_elem)
    return tbl_elem


def insert_figure_after(doc, anchor, image_path, caption_text, width_in):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width_in))
    img_elem = img_p._p
    img_elem.getparent().remove(img_elem)

    cap_p = doc.add_paragraph()
    cap_run = cap_p.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_elem = cap_p._p
    cap_elem.getparent().remove(cap_elem)

    anchor_elem = _elem_of(anchor)
    anchor_elem.addnext(cap_elem)
    anchor_elem.addnext(img_elem)
    return cap_elem


# =======================================================================
# content text
# =======================================================================
INTRO_TEXT = (
    "Round-12 point 1 (Cauchy-stress fixed-region QoI, final checkpoints): the advisor "
    "asked for one final accuracy-versus-resolution table -- for EACH resolution, FEM "
    "vs. operator vs. the same fine reference, in displacement L2, H1/energy norm, "
    "reaction force, and stress -- with Cauchy stress (not PK1) as the main engineering "
    "quantity, explicitly rejecting the bare pointwise maximum as the primary QoI "
    "(singular/mesh-dependent at corners even for FEM) in favour of a FIXED physical "
    "region around the stress concentration, reporting a robust local statistic (a "
    "weighted average or 99th percentile) with the true max still available separately, "
    "the region fixed in physical space across every resolution. All new code "
    "(fixed-region selection, the Cauchy push-forward from the already-validated PK1 "
    "stress, and the weighted-percentile/top-fraction statistics) was verified on CPU "
    "before any GPU time was spent (5 checks, all passed). Below: two tables per case, "
    "all six cases, all sixteen resolutions tested (N=3..49) -- the established "
    "accuracy QoIs, then the new Cauchy-stress QoI. Scope, stated plainly: these "
    "resolutions are the LOW_N range already established for this project's own "
    "coarsest-suitable-mesh crossover analysis above, against a fine reference at "
    "N=201 -- this does NOT extend to N=1401, which would need a fresh, expensive "
    "fine_N=2236 reference for the five cases that have never had one, a separate, "
    "bigger ask that should be scoped on its own rather than silently bundled in here."
)

GAP_NOTE = (
    "A genuine, pre-existing gap, disclosed rather than hidden: the operator's own "
    "classical accuracy QoIs (L2, H1, energy, reaction) were never evaluated below "
    "N=13 for any case before this round -- its own zero-shot-validated range has "
    "always started there. B1xNeo-Hookean's own \"Op.\" columns below are complete "
    "down to N=3 only because a dedicated follow-up run (2026-09-18, 46s of real GPU "
    "compute) closed that gap specifically for the flagship case; the other five "
    "cases still show \"n/a\" for the operator at N=3,4,5,6,9,11, an honest, "
    "unfilled gap rather than a guess. The Cauchy-stress metric, by contrast, is "
    "complete for every case at every resolution, since it was computed fresh in "
    "this round's own sweep regardless of any older cache."
)

CLOSING_DISCUSSION = (
    "Two findings worth stating plainly. First, the true max is markedly noisier and "
    "slower-converging than the region average for BOTH methods and every case -- "
    "B1xNeo-Hookean's own FEM region average falls to 0.30% by N=49 (displacement L2 "
    "is already down to 0.05% there), while its FEM true max is still 62.8% at N=3 "
    "and only reaches 26.5% at N=49 -- a real, data-grounded illustration of exactly "
    "why the advisor asked to avoid the bare pointwise maximum as the primary design "
    "QoI: even FEM's own true max, with no operator involved at all, converges far "
    "more slowly than the region-averaged statistic he asked to use instead. Second, "
    "the operator's own region-averaged stress error is consistently worse than its "
    "displacement L2 error at matched resolution (e.g. B1xMooney-Rivlin at N=13: L2 "
    "7.39% vs. region average 10.69%) -- stress is a genuinely harder, noisier target "
    "for the operator than displacement, not merely a rescaled version of the same "
    "accuracy. B2xNeo-Hookean is the extreme case: its operator's own region-averaged "
    "stress error reaches 245.59% at N=3 and stays in the 3.2%-245.6% range across "
    "the whole sweep, never settling anywhere close to FEM's own sub-1% region "
    "average at the same resolutions -- confirming that displacement accuracy alone "
    "should not be treated as a stand-in for stress accuracy."
)

FLAGSHIP_FIGURE_CAPTION = (
    f"Round-12 Figure A. B1 x Neo-Hookean, region-average vs. true-max Cauchy stress "
    f"error, FEM and operator, both vs. the same fine reference (Table "
    f"{TABLE_ID[('B1_neo_hookean', 'cauchy')]})."
)


def classical_caption(case_id):
    tid = TABLE_ID[(case_id, 'classical')]
    label = CASE_LABELS[case_id]
    extra = ''
    if case_id == 'B1_neo_hookean':
        extra = (' The operator\'s own N=3..11 cells here come from a dedicated '
                 'follow-up run against the same final checkpoint (see the gap note above).')
    elif case_id.startswith('B2'):
        extra = ' Reaction cells are n/a for both methods (no reaction resultant defined for the B2 ring geometry).'
    return (
        f"Table {tid}. {label}: displacement L2, H1 semi-norm, tangent-energy norm, "
        f"and reaction-force error, FEM vs. operator, all sixteen resolutions tested, "
        f"same fine reference (N=201).{extra}"
    )


def cauchy_caption(case_id):
    tid = TABLE_ID[(case_id, 'cauchy')]
    label = CASE_LABELS[case_id]
    return (
        f"Table {tid}. {label}: fixed-region Cauchy-stress error (region-weighted "
        f"average, 99th percentile, true max), FEM vs. operator, all sixteen "
        f"resolutions tested, same fine reference and fixed-region convention as "
        f"the L2/H1/energy/reaction table above."
    )


# =======================================================================
# apply to the Report
# =======================================================================
doc = Document(REPORT)
paras = list(doc.paragraphs)

START_TEXT = paras[[i for i, p in enumerate(paras)
                     if p.text.strip().startswith('Round-12 point 1 (Cauchy-stress')][0]].text.strip()
END_TEXT = (
    "The operator's own region-averaged stress error is consistently worse than its "
    "displacement L2 error at this same resolution (e.g. B1xMooney-Rivlin: L2 7.39% vs. "
    "region average 10.69%) -- stress is a genuinely harder, noisier target for the "
    "operator than displacement, not merely a rescaled version of the same accuracy. "
    "B2xNeo-Hookean is the extreme case: its operator's own region-averaged stress error "
    "reaches 245.59% at N=3 and stays in the 3.2%-245.6% range across the whole LOW_N "
    "sweep tested, never settling anywhere close to FEM's own sub-1% region average at "
    "the same resolutions -- a real finding, not a display artefact, confirming that "
    "displacement accuracy alone should not be treated as a stand-in for stress accuracy "
    "for this case."
)

anchor = remove_range_return_anchor(doc, START_TEXT, END_TEXT)

p = insert_paragraph_after(doc, anchor, INTRO_TEXT)
p = insert_paragraph_after(doc, p, GAP_NOTE)

for case_id in CASE_ORDER:
    p = insert_paragraph_after(doc, p, CASE_LABELS[case_id] + ':', bold=True)
    p = insert_paragraph_after(doc, p, classical_caption(case_id))
    p = insert_table_after(doc, p, CLASSICAL_HEADER, classical_rows(case_id))
    p = insert_paragraph_after(doc, p, cauchy_caption(case_id))
    p = insert_table_after(doc, p, CAUCHY_HEADER, cauchy_rows(case_id))
    if case_id == 'B1_neo_hookean':
        p = insert_figure_after(doc, p, os.path.join(FIG, 'fig_round12_cauchy_flagship.png'),
                                 FLAGSHIP_FIGURE_CAPTION, width_in=6.0)

p = insert_paragraph_after(doc, p, CLOSING_DISCUSSION)

doc.save(REPORT)
print('Saved', REPORT)

check = Document(REPORT)
print('paragraphs:', len(check.paragraphs), 'tables:', len(check.tables), 'images:', len(check.inline_shapes))
