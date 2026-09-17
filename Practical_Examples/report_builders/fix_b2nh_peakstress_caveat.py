"""Real accuracy bug caught during the pre-send audit Omar asked for:
Table R10-6 / Table 18-R10k (peak-stress QoI, 5 cases) lists
B2xNeo-Hookean's disp_rel_L2 as 46.37% -- which matches its ORIGINAL
(pre-retrain) checkpoint (Table 18-R10l / Table R10-2's own "OLD"
column), not the now-published retrained checkpoint's 36.05%. That is
correct and self-consistent (the peak-stress number in the same row is
ALSO from the original checkpoint, since task #22 predates the retrain)
-- but unlike the other three caveated rows (B1xMooney-Rivlin,
B1xArruda-Boyce, B2xMooney-Rivlin), this row's own peak-stress cell was
never annotated "(original checkpoint)", and the surrounding paragraph
still called B2xNeo-Hookean's retrain "still mid-training" -- stale
now that it has finished and been published elsewhere in the same
document. Left as-is, a reader would reasonably assume this row's
46.37%/49.48% pair reflects the NEW checkpoint everywhere else in the
document, since the document elsewhere reports B2xNeo-Hookean's own
36.05% at N=1401.

Fixes both real documents: adds the missing "(original checkpoint)"
annotation to B2xNeo-Hookean's own peak-stress cell, and corrects the
caveat paragraph to say four rows (not three) share the caveat, with
only B2xArruda-Boyce (no retrain planned) genuinely caveat-free.
"""
import os

from docx import Document

DELIV = '/home/user/OMAR/advisor_feedback'

REPORT_SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17e.docx')
REPORT_DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-17f.docx')
SUMMARY_SRC = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-17f.docx')
SUMMARY_DST = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-17g.docx')


def find_para_exact(doc, original, text):
    hits = [p for p in original if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def replace_paragraph_text(para, new_text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(new_text)
    return para


def fix_table_cell(doc, header, row_label, old_val, new_val):
    for t in doc.tables:
        h = [c.text for c in t.rows[0].cells]
        if h == header:
            for r in t.rows[1:]:
                if r.cells[0].text == row_label:
                    cell = r.cells[2]
                    assert cell.text == old_val, f'expected {old_val!r}, found {cell.text!r}'
                    for p in cell.paragraphs:
                        for run in list(p.runs):
                            run._r.getparent().remove(run._r)
                    cell.paragraphs[0].add_run(new_val)
                    return
    raise AssertionError(f'table with header {header} / row {row_label} not found')


# ------------------------------------------------------------------
# REPORT
# ------------------------------------------------------------------
doc = Document(REPORT_SRC)
ORIGINAL = list(doc.paragraphs)

# Same pre-send audit also found the Report missing explicit "round-11
# point N" labels for points 1 and 3 (points 2 and 4 already had them) --
# fixed here for consistency, same as the Summary's own fix.
p_protocol = find_para_exact(doc, ORIGINAL,
    "The exact multi-resolution training protocol, since the advisor asked "
    "for it precisely rather than by description: weighting across the four "
    "resolutions is EQUAL BY CONSTRUCTION, not importance-weighted. 400 "
    "training samples and 100 validation samples are generated per "
    "resolution, an identical count at all four. Every epoch rebuilds a fresh "
    "batch plan: each individual batch is homogeneous in N (one forward pass "
    "shares one mesh/quadrature tensor, so a batch cannot mix resolutions), "
    "but the full list of batches spanning all four resolutions is then "
    "shuffled together, so batches from different resolutions interleave "
    "randomly through the epoch rather than training block by block, one "
    "resolution at a time. The loss itself is the plain per-batch mean of the "
    "potential energy (Pi = U - W); it carries no explicit cross-resolution "
    "weighting term, so the only sense in which resolutions are \"weighted\" "
    "is via their equal sample counts. Validation error is reported as the "
    "unweighted mean of the four resolutions' own per-resolution mean errors, "
    "regardless of how difficult each one is. N=21 and N=33 were kept because "
    "they are the base checkpoint's own original training resolutions; N=101 "
    "and N=201 were added because the diagnostic accuracy-degradation sweep "
    "(Table 18-R10a) had already measured real, growing error at exactly "
    "those two points (15.0% at N=101, 22.3% at N=201) before any retraining "
    "ran -- the choice is staged evidence, not an arbitrary spread, and "
    "deliberately stops short of adding even wider (401/701) resolutions "
    "until this moderate-cost fix was shown to help."
)
replace_paragraph_text(p_protocol,
    "Round-11 point 3 (the exact multi-resolution training protocol, since "
    "the advisor asked for it precisely rather than by description): "
    "weighting across the four resolutions is EQUAL BY CONSTRUCTION, not "
    "importance-weighted. 400 training samples and 100 validation samples are "
    "generated per resolution, an identical count at all four. Every epoch "
    "rebuilds a fresh batch plan: each individual batch is homogeneous in N "
    "(one forward pass shares one mesh/quadrature tensor, so a batch cannot "
    "mix resolutions), but the full list of batches spanning all four "
    "resolutions is then shuffled together, so batches from different "
    "resolutions interleave randomly through the epoch rather than training "
    "block by block, one resolution at a time. The loss itself is the plain "
    "per-batch mean of the potential energy (Pi = U - W); it carries no "
    "explicit cross-resolution weighting term, so the only sense in which "
    "resolutions are \"weighted\" is via their equal sample counts. "
    "Validation error is reported as the unweighted mean of the four "
    "resolutions' own per-resolution mean errors, regardless of how difficult "
    "each one is. N=21 and N=33 were kept because they are the base "
    "checkpoint's own original training resolutions; N=101 and N=201 were "
    "added because the diagnostic accuracy-degradation sweep (Table 18-R10a) "
    "had already measured real, growing error at exactly those two points "
    "(15.0% at N=101, 22.3% at N=201) before any retraining ran -- the choice "
    "is staged evidence, not an arbitrary spread, and deliberately stops "
    "short of adding even wider (401/701) resolutions until this "
    "moderate-cost fix was shown to help."
)

p_breakeven = find_para_exact(doc, ORIGINAL,
    "The advisor's round-11 feedback asked to keep BOTH break-even "
    "comparisons in the report, expecting the second kind -- both methods "
    "evaluated at the SAME resolution, N=1401, rather than at each method's "
    "own accuracy-matched mesh -- to look substantially more favourable to "
    "the operator. Table 18-R10e' (below) gives exactly that: the operator's "
    "own default eager fp32 forward pass at N=1401 against torch-fem's own "
    "real GPU solve time at the SAME N=1401, for all six (geometry, material) "
    "cases."
)
replace_paragraph_text(p_breakeven,
    "Round-11 point 1 (keep BOTH break-even comparisons): the advisor asked "
    "to keep BOTH break-even comparisons in the report, expecting the second "
    "kind -- both methods evaluated at the SAME resolution, N=1401, rather "
    "than at each method's own accuracy-matched mesh -- to look substantially "
    "more favourable to the operator. Table 18-R10e' (below) gives exactly "
    "that: the operator's own default eager fp32 forward pass at N=1401 "
    "against torch-fem's own real GPU solve time at the SAME N=1401, for all "
    "six (geometry, material) cases."
)

p = find_para_exact(doc, ORIGINAL,
    "Task #22 also asked for the peak-stress QoI (fixed physical location, "
    "converged from a fine reference -- the same convention as B1xNeo-Hookean's "
    "own peak-stress discussion above) at N=1401 for the remaining five cases. "
    "Table 18-R10k. A checkpoint-version caveat applies to three of the five "
    "rows: B1xMooney-Rivlin, B1xArruda-Boyce and B2xMooney-Rivlin were measured "
    "here against their ORIGINAL (pre-retrain) checkpoints -- their "
    "disp_rel_L2 values match exactly the \"OLD\" column already published in "
    "Tables 18-R10f/g/h -- so the peak-stress numbers below have not been "
    "recomputed against this week's own newly retrained checkpoints for those "
    "three cases. B2xNeo-Hookean (its own first multi-res retrain still "
    "mid-training) and B2xArruda-Boyce (no retrain planned) have only ever had "
    "one checkpoint each, so no such caveat applies to those two rows."
)
replace_paragraph_text(p,
    "Round-11 point 2, part A (extend the N=1401 QoI analysis to all "
    "remaining cases): task #22 also asked for the peak-stress QoI (fixed "
    "physical location, converged from a fine reference -- the same "
    "convention as B1xNeo-Hookean's own peak-stress discussion above) at "
    "N=1401 for the remaining five cases. Table 18-R10k. A checkpoint-version "
    "caveat applies to FOUR of the five "
    "rows: B1xMooney-Rivlin, B1xArruda-Boyce, B2xMooney-Rivlin and "
    "B2xNeo-Hookean were measured here against their ORIGINAL (pre-retrain) "
    "checkpoints -- their disp_rel_L2 values match exactly the \"OLD\" column "
    "already published in Tables 18-R10f/g/h/l -- so the peak-stress numbers "
    "below have not been recomputed against this week's own newly retrained "
    "checkpoints for those four cases. B2xArruda-Boyce (no retrain planned) "
    "has only ever had one checkpoint, so no such caveat applies to that row."
)

fix_table_cell(doc, ['Case', 'disp_rel_L2 @N=1401', 'peak_stress_rel_err @N=1401'],
               'B2 x Neo-Hookean', '49.48%', '49.48% (original checkpoint)')

doc.save(REPORT_DST)
print('Saved', REPORT_DST)

# ------------------------------------------------------------------
# SUMMARY
# ------------------------------------------------------------------
doc = Document(SUMMARY_SRC)
ORIGINAL = list(doc.paragraphs)

p = find_para_exact(doc, ORIGINAL,
    "Round-11 point 2, part A (extend the N=1401 QoI analysis to all remaining "
    "cases): task #22's peak-stress QoI, the five cases beyond B1xNeo-Hookean "
    "(Table R10-6). Caveat: B1xMooney-Rivlin, B1xArruda-Boyce and "
    "B2xMooney-Rivlin here reflect their ORIGINAL pre-retrain checkpoints "
    "(their disp_rel_L2 matches the \"OLD\" column in Table R10-2 exactly) -- "
    "the peak-stress metric has not been recomputed against this week's new "
    "retrained checkpoints for those three. B2xNeo-Hookean and B2xArruda-Boyce "
    "have only one checkpoint each, no caveat."
)
replace_paragraph_text(p,
    "Round-11 point 2, part A (extend the N=1401 QoI analysis to all remaining "
    "cases): task #22's peak-stress QoI, the five cases beyond B1xNeo-Hookean "
    "(Table R10-6). Caveat: B1xMooney-Rivlin, B1xArruda-Boyce, "
    "B2xMooney-Rivlin and B2xNeo-Hookean here reflect their ORIGINAL "
    "pre-retrain checkpoints (their disp_rel_L2 matches the \"OLD\" column in "
    "Table R10-2 exactly) -- the peak-stress metric has not been recomputed "
    "against this week's new retrained checkpoints for those four. "
    "B2xArruda-Boyce has only one checkpoint, no caveat."
)

fix_table_cell(doc, ['Case', 'disp_rel_L2 @N1401', 'peak_stress_rel_err @N1401'],
               'B2 x Neo-Hookean', '49.48%', '49.48% (orig. checkpoint)')

doc.save(SUMMARY_DST)
print('Saved', SUMMARY_DST)
