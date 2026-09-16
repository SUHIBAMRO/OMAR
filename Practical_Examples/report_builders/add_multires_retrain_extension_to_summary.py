"""Companion to add_multires_retrain_extension_to_report.py -- same new
content, condensed to the Work Summary document's own terser style and
table/figure naming convention (Table Rxx-N, Figure Rxx-X, not the full
Report's Table 18-R10x scheme). See that script's own docstring for the
full rationale and exact number provenance; not repeated here.

Inserted right after the Summary's own paragraph describing the peak-
stress flip for the B1xNeo-Hookean retrain (its own "Table R10-1"
paragraph group) and before "Point 2 -- batch size and throughput...",
the same natural continuation point used in the Report."""
import copy
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

FIG = '/home/user/OMAR/Practical_Examples/report_builders/figures'
DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-15.docx')
DST = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-16.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_table_after(anchor_para, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    pr = doc.tables[0]._tbl.find(qn('w:tblPr'))
    if pr is not None:
        old = tbl._tbl.find(qn('w:tblPr'))
        if old is not None:
            tbl._tbl.remove(old)
        tbl._tbl.insert(0, copy.deepcopy(pr))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    tbl._tbl.getparent().remove(tbl._tbl)
    anchor_para._p.addnext(tbl._tbl)
    return tbl


def insert_after_table(table, style_para, text):
    new_p_el = copy.deepcopy(style_para._p)
    table._tbl.addnext(new_p_el)
    np = Paragraph(new_p_el, style_para._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_figure_after(anchor_para, image_path, caption_text, width_in=5.5):
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(image_path, width=Inches(width_in))
    img_elem = img_p._p
    img_elem.getparent().remove(img_elem)

    cap_p = doc.add_paragraph()
    cap_run = cap_p.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(9)
    cap_elem = cap_p._p
    cap_elem.getparent().remove(cap_elem)

    anchor_para._p.addnext(cap_elem)
    anchor_para._p.addnext(img_elem)
    return Paragraph(cap_elem, anchor_para._parent)


anchor = find_para_exact(
    'The retrained checkpoint also flips the peak-stress finding: its peak-stress '
    'accuracy (42-70%, improved at every N) is now BETTER than anything FEM achieves '
    'in the whole low-N range (62% at best) for every operator resolution N>=29 -- FEM '
    'would need N>49 (untested) to match it. The "coarsest suitable FEM" is now bound '
    'by tangent energy, not peak stress, and dropped from N=17-45 (original '
    'checkpoint) to a near-degenerate N=9-17 (retrained checkpoint) for almost every '
    'resolution. Same domain-corner caveat still applies to peak stress specifically '
    '-- the improvement itself is real and independently verified, not a '
    'training-validation artifact.'
)

p1 = insert_after(anchor,
    'The same fix (retrain on N=21,33,101,201 instead of N=21,33 only) was since '
    'extended to every other case with the same degradation signature. Protocol, '
    'exactly as the advisor asked: weighting is equal by construction (400 train / '
    '100 val samples per resolution, identical count at all four); each batch is '
    'homogeneous in one resolution, but the full batch list spanning all four is '
    'shuffled together so different resolutions interleave randomly through the '
    'epoch; the loss carries no explicit cross-resolution weight. N=101/201 were '
    'added because they were exactly where the degradation sweep had already '
    'measured real error (15.0%/22.3%) before any retraining ran. Two methods '
    'changes made the later retrains cheaper: TF32 training (verified safe, 2.08x '
    'at N=201, the resolution it actually helps), and a real bug found and fixed '
    'this session -- the GPU fast data-generation path had only ever been wired up '
    'for B1, so B2\'s own --fast_solver/--nsteps did nothing at all until now; every '
    'earlier B2 job silently ran on the slow CPU solver regardless of its own launch '
    'command.')

p2 = insert_after(p1,
    'Results, N=1401, old vs. new checkpoint (Table R10-2): B1 x Mooney-Rivlin '
    '39.20% -> 15.04% (~62% relative reduction); B1 x Arruda-Boyce 45.62% -> 25.05% '
    '(~45%, and its own real GPU run is reassuring evidence that a shared '
    'chain-locking-clamp exposure with this project\'s own torch-fem Arruda-Boyce '
    'runs does not bite in practice here -- all 16 resolutions, both checkpoints, '
    'converged cleanly); B2 x Mooney-Rivlin 49.47% -> 13.10% (~73.5%, the largest '
    'reduction so far, and the clearest case of a genuinely FLAT error across '
    'N=37-1401 rather than merely a smaller one at N=1401). B2 x Neo-Hookean\'s own '
    'retrain and the advisor\'s round-11 direct-N1401 ablation were both still '
    'mid-training as of this writing and are not included below.')

header = ['Case', 'Old (N=1401)', 'New (N=1401)', 'Relative reduction']
t = insert_table_after(p2, header, [
    ['B1 x Neo-Hookean', '44.65%', '5.85%', '~87% (7.6x)'],
    ['B1 x Mooney-Rivlin', '39.20%', '15.04%', '~62%'],
    ['B1 x Arruda-Boyce', '45.62%', '25.05%', '~45%'],
    ['B2 x Mooney-Rivlin', '49.47%', '13.10%', '~73.5%'],
])
p3 = insert_after_table(t, p2,
    'Table R10-2. Multi-resolution retrain fix, every completed case, disp_rel_L2 at '
    'N=1401.')

p4 = insert_figure_after(p3,
    os.path.join(FIG, 'fig_multires_retrain_summary.png'),
    'Figure R10-B. Multi-resolution retrain fix, old vs. new checkpoint, all '
    'completed cases (Table R10-2).')

doc.save(DST)
print('Saved', DST)
