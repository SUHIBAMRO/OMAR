"""Item #11, Summary counterpart: extend section 11 (Verification against
an analytic solution / MMS, Tables 22-24) beyond Neo-Hookean to
Mooney-Rivlin and Arruda-Boyce, into the REAL, already-updated
deliverable (PFEM_Work_Summary_updated_2026-09-07.docx). Mirrors
add_item_11_mms_materials_report.py's data and table numbering (22a,
22b, 23a) but in the Summary's own terser voice.
"""
import copy
import json
import os

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

PF = '/home/user/OMAR/Practical_Examples/omar_pfem'
DELIV = '/tmp/claude-0/-home-user/64d7c4d8-d5f0-5686-a58f-aa87abfd4ba4/scratchpad/deliverables'
SRC = os.path.join(DELIV, 'PFEM_Work_Summary_updated_2026-09-07.docx')
DST = os.path.join(DELIV, 'PFEM_Work_Summary_updated_2026-09-07b.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_table_after(anchor_para, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    pr = doc.tables[0]._tbl.find(qn('w:tblPr'))
    if pr is not None:
        old = tbl._tbl.find(qn('w:tblPr'))
        if old is not None:
            tbl._tbl.remove(old)
        tbl._tbl.insert(0, copy.deepcopy(pr))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    tbl._tbl.getparent().remove(tbl._tbl)
    anchor_para._p.addnext(tbl._tbl)
    return tbl


def insert_after_table(table, style_para, text):
    new_p_el = copy.deepcopy(style_para._p)
    table._tbl.addnext(new_p_el)
    np = Paragraph(new_p_el, style_para._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def load_rows(material):
    d = json.load(open(os.path.join(PF, 'point9_results', f'mms_B1_{material}.json')))
    return d['rows'], d['convergence_rates'], d['rate_check']


def fmt(v):
    return f'{v:.3e}'


def table_rows_for(rows):
    return [[r['order'], r['N'], f"{r['n_dof']:,}", fmt(r['L2_rel']), fmt(r['H1_semi_rel']),
             fmt(r['stress_rel_L2']), fmt(r['energy_rel'])] for r in rows]


def rate_rows_for(material_label, rates):
    out = []
    for order in ('Q4', 'Q9'):
        for norm_key, norm_label, theory in [
            ('L2', 'L2', 2 if order == 'Q4' else 3),
            ('H1_semi', 'H1 semi-norm', 1 if order == 'Q4' else 2),
            ('stress', 'Stress', 1 if order == 'Q4' else 2),
            ('energy', 'Energy', 2 if order == 'Q4' else 4),
        ]:
            r = rates[order][norm_key]
            pw = ', '.join(f'{v:.2f}' for v in r['pairwise'])
            out.append([material_label, order, norm_label, f"{r['rate']:.2f}", theory, pw])
    return out


mr_rows, mr_rates, mr_check = load_rows('mooney_rivlin')
ab_rows, ab_rates, ab_check = load_rows('arruda_boyce')
assert mr_check == {'Q4': 'as expected', 'Q9': 'as expected'}
assert ab_check == {'Q4': 'as expected', 'Q9': 'as expected'}

HEADER = ['Order', 'N', 'DOF', 'L2', 'H1 semi-norm', 'Stress', 'Energy']

anchor = find_para_exact(
    'operator / Q4 = 2.42× in L2 at this mesh. The finding is that the four norms '
    'disagree: 1.03× in H1 and 1.03× in stress effectively at the Q4 optimum against '
    '2.42× in L2 and 3.11× in energy. That inverts the usual ordering. The loss is '
    'built from the deformation gradient, so strain and stress are what it constrains '
    'hardest, and the displacement is pinned only through them; the same inversion '
    'appears in an independent N = 9 CPU run (1.35× against 4.71×). For a '
    'physics-informed operator an L2 displacement error overstates how wrong the '
    'mechanics are.'
)

p1 = insert_after(
    anchor,
    'Tables 22-24 verify Neo-Hookean only. The same manufactured field and B1 mesh '
    'sequence were repeated for Mooney-Rivlin and Arruda-Boyce to confirm the check '
    'is not accidental to one energy functional.'
)

t22a = insert_table_after(p1, HEADER, table_rows_for(mr_rows))
p2 = insert_after_table(
    t22a, p1, 'Table 22a. B1 x Mooney-Rivlin against the manufactured solution.')

t22b = insert_table_after(p2, HEADER, table_rows_for(ab_rows))
p3 = insert_after_table(
    t22b, p2, 'Table 22b. B1 x Arruda-Boyce against the manufactured solution.')

rate_header = ['Material', 'Order', 'Norm', 'Observed rate', 'Theory', 'Pairwise']
rate_rows = rate_rows_for('Mooney-Rivlin', mr_rates) + rate_rows_for('Arruda-Boyce', ab_rates)
t23a = insert_table_after(p3, rate_header, rate_rows)
p4 = insert_after_table(
    t23a, p3,
    'Table 23a. Observed rates for both materials, fitted the same way as Table 23. '
    'All sixteen land on theory -- the verification holds for every material this '
    'report trains an operator on, not only the one in Tables 22-24.'
)

doc.save(DST)
print('Saved', DST)
