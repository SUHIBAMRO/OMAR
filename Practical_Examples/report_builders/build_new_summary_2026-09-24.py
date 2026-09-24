"""Non-cumulative Summary (Omar's own 2026-09-18 convention: Summary
covers ONLY the round just completed, copied verbatim out of the now-
updated cumulative Report -- see build_new_summary_2026-09-18.py for the
original rationale/mechanism, reused unchanged here).

This round: the full GPU mesh-convergence results for both 3D candidates
(B3 sharper-groove and rigid-shim laminated bushing), Prof. Rabczuk's
final candidate decision (B3 chosen), and the decided VINO/Transolver
dataset-variation scope. Slices Section 11 in full (11.1 through 11.5)
out of PFEM_Transolver_Report_2026-09-24b.docx -- the round-15 and
round-16 report-builder scripts' own docstrings hold the internal record
of what changed across drafts; this file's own text is what actually
goes to the advisor.
"""
import copy
import os

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

DELIV = '/home/user/OMAR/advisor_feedback'
REPORT = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-24b.docx')
OUT = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-24.docx')

src = Document(REPORT)


def body_children(doc):
    return list(doc.element.body)


def elem_text(el):
    if el.tag != qn('w:p'):
        return None
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_index(children, text):
    matches = [i for i, el in enumerate(children)
               if elem_text(el) is not None and elem_text(el).strip() == text]
    assert len(matches) == 1, f'{len(matches)} matches for {text!r}'
    return matches[0]


def slice_by_anchor_text(doc, start_text, end_text):
    children = body_children(doc)
    i0 = find_index(children, start_text)
    i1 = find_index(children, end_text)
    assert i1 >= i0
    return children[i0:i1 + 1]


def append_heading(dst_doc, text):
    from docx.shared import Pt
    p = dst_doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)


def copy_elements(dst_doc, src_doc, elements):
    body = dst_doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    for el in elements:
        new_el = copy.deepcopy(el)
        for blip in new_el.iter(qn('a:blip')):
            rId = blip.get(qn('r:embed'))
            if not rId:
                continue
            image_part = src_doc.part.related_parts[rId]
            new_rId = dst_doc.part.relate_to(image_part, RT.IMAGE)
            blip.set(qn('r:embed'), new_rId)
        if sect_pr is not None:
            sect_pr.addprevious(new_el)
        else:
            body.append(new_el)


dst = Document()
if src.tables:
    src_style = src.tables[0].style
    try:
        dst.styles.add_style(src_style.name, src_style.type)
    except ValueError:
        pass

append_heading(dst, "PFEM / Transolver Work Summary -- Third 3D Candidate, "
                     "Final Decision")
p = dst.add_paragraph()
p.add_run(
    "This round covers the completed GPU mesh-convergence study for both "
    "candidates prepared for item 4 of a previous round (a new, harder, "
    "more realistic 3D example): B3 with a sharper groove, and a laminated "
    "rubber-mount bushing with rigid steel shims. Following review of both "
    "candidates' complete results, Prof. Rabczuk selected B3 for the final "
    "FEM-versus-operator comparison; this also covers the resulting scope "
    "of which geometry/material/loading parameters will vary in the "
    "training dataset. As with previous rounds, everything below is copied "
    "verbatim (same tables, same text) out of the now-updated cumulative "
    "Report (PFEM_Transolver_Report_2026-09-24b.docx) -- nothing "
    "summarized or re-derived. No dataset generation or operator training "
    "has started."
).italic = True

elements = slice_by_anchor_text(
    src,
    "11. Third Benchmark Candidate: a 3D Rocking Rubber-Mount Bushing "
    "(Two Design Variants Compared)",
    "No dataset generation or neural-operator training has started. After "
    "the FEM time/memory table is complete and the sampling ranges above "
    "are set, the remaining numerical steps are dataset generation, "
    "operator training, and the FEM-versus-operator comparison on "
    "displacement, reaction, energy, and regional Cauchy-stress QoIs, "
    "exactly as done for B1 and B2 -- none of this is implied to be "
    "complete by anything above.",
)
copy_elements(dst, src, elements)

dst.save(OUT)
print('Saved', OUT)

check = Document(OUT)
print('paragraphs:', len(check.paragraphs))
print('tables:', len(check.tables))
print('images:', len(check.inline_shapes))
