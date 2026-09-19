"""Round-13 item 2 (Timon's newest email): "use your most efficient
validated GPU-native FEM solver as the primary timing baseline."

Investigated via direct code inspection (not assumption, 2026-09-19):
Table 18-R10e's own "finite-element solver@N=11" measurement already
comes from `gpu_fem_benchmark.py` -> `gpu_fem_solver.py` -- OUR OWN
GPU-native Total-Lagrangian Newton solver, built earlier per the
advisor's own request for a GPU-native comparison, completely
independent of torch-fem (which is used only in the SEPARATE
resolution-matched table, 18-R10e', matching the advisor's own request
to "keep the same-N comparison separately"). So item 2 needs no
computational change -- only making this explicit in the Report text,
since the cell that produces this number (cell_break_even_accuracy_
matched.py) only names it "torch-fem" in its own informal print()
statements, never in the actual Report wording.

Pure text edit: two paragraphs (the Table 18-R10e discussion and its own
caption) get one clarifying sentence/clause each. No new paragraphs,
tables, or images.
"""
import os

from docx import Document

DELIV = '/home/user/OMAR/advisor_feedback'
REPORT_SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-19.docx')
REPORT_DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-19b.docx')

OLD_DISCUSSION = (
    "Point 5 asked for an updated break-even analysis. Using the retrained checkpoint's "
    "own coarsest-suitable finite-element mesh at N=1401 (N=11, bound by the "
    "tangent-energy norm per the crossover above) against the operator's own real "
    "inference cost there, and the real training wall-clock for the retrained checkpoint "
    "(41,881 s, about 11.6 hours, confirmed above): in its default eager fp32 deployment "
    "mode, the operator never recovers its own training cost against this baseline -- the "
    "accuracy-matched finite-element mesh is already cheaper per sample (1,625.6 ms) than "
    "the operator's default forward pass (2,292.1 ms). With the torch.compile plus TF32 "
    "optimization above (394.0 ms/sample), the operator is instead 4.13x faster per "
    "sample and recovers its training cost after 34,005 solved problems, about 3.72 "
    "GPU-hours of inference against the 11.6 GPU-hours spent training. This is the only "
    "point in the whole round-10 investigation where the operator's default, unoptimized "
    "deployment mode does not come out ahead; its economic case against a genuinely "
    "accuracy-matched baseline rests on deploying it with the optimizations found under "
    "point 3, not on its default settings."
)
NEW_DISCUSSION = OLD_DISCUSSION + (
    " (The finite-element side of this comparison is our own GPU-native Total-Lagrangian "
    "Newton solver, gpu_fem_solver.py -- built earlier per the advisor's own request for a "
    "GPU-native comparison, and already the primary timing baseline used here; torch-fem "
    "appears only in the separate resolution-matched comparison below, Table 18-R10e', "
    "which intentionally asks a different question -- same-N, not accuracy-matched -- per "
    "the advisor's own request to keep the two comparisons apart.)"
)

OLD_CAPTION = (
    "Table 18-R10e. Accuracy-matched break-even: operator@N=1401 vs. finite-element "
    "solver@N=11 (its coarsest suitable mesh for the retrained checkpoint). Distinct from "
    "the matched-resolution/matched-batch-size break-even of Section 8.4: this comparison "
    "intentionally runs the two methods at different N, matched by accuracy rather than by "
    "mesh."
)
NEW_CAPTION = (
    "Table 18-R10e. Accuracy-matched break-even: operator@N=1401 vs. our own GPU-native "
    "finite-element solver (gpu_fem_solver.py) @N=11 (its coarsest suitable mesh for the "
    "retrained checkpoint). Distinct from the matched-resolution/matched-batch-size "
    "break-even of Section 8.4: this comparison intentionally runs the two methods at "
    "different N, matched by accuracy rather than by mesh."
)


def replace_paragraph_text(para, new_text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    para.add_run(new_text)


doc = Document(REPORT_SRC)
paras = list(doc.paragraphs)
before_p, before_t, before_i = len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes)

disc_hits = [p for p in paras if p.text.strip() == OLD_DISCUSSION]
assert len(disc_hits) == 1, f'{len(disc_hits)} matches for discussion paragraph'
replace_paragraph_text(disc_hits[0], NEW_DISCUSSION)

cap_hits = [p for p in paras if p.text.strip() == OLD_CAPTION]
assert len(cap_hits) == 1, f'{len(cap_hits)} matches for caption paragraph'
replace_paragraph_text(cap_hits[0], NEW_CAPTION)

doc.save(REPORT_DST)
print('Saved', REPORT_DST)

check = Document(REPORT_DST)
after_p, after_t, after_i = len(check.paragraphs), len(check.tables), len(check.inline_shapes)
print(f'paragraphs: {before_p}->{after_p}, tables: {before_t}->{after_t}, images: {before_i}->{after_i}')
assert (before_p, before_t, before_i) == (after_p, after_t, after_i), 'structure changed -- pure text edit expected'
