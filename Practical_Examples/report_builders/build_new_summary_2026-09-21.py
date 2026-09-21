"""Round-14 non-cumulative Summary (Omar's own 2026-09-18 convention:
Summary covers ONLY the round just completed, copied verbatim out of the
now-updated cumulative Report -- see build_new_summary_2026-09-18.py for
the original rationale/mechanism, reused unchanged here).

Round 14 = Timon's item 4 (a new, harder 3D example): first status
report on both candidates (B3 rubber-mount bushing, tire sector),
Omar's own detailed 11-point technical review of both, the real GPU
mesh-convergence result that followed, and two genuine field-error fix
attempts that were tried, tested, and reverted.
"""
import copy
import os

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

DELIV = '/home/user/OMAR/advisor_feedback'
REPORT = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-21.docx')
OUT = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-21.docx')

src = Document(REPORT)


def body_children(doc):
    return list(doc.element.body)


def elem_text(el):
    if el.tag != qn('w:p'):
        return None
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_index(children, text):
    matches = [i for i, el in enumerate(children)
               if elem_text(el) is not None and elem_text(el).strip() == text]
    assert len(matches) == 1, f'{len(matches)} matches for {text!r}'
    return matches[0]


def slice_by_anchor_text(doc, start_text, end_text):
    children = body_children(doc)
    i0 = find_index(children, start_text)
    i1 = find_index(children, end_text)
    assert i1 >= i0
    return children[i0:i1 + 1]


def append_heading(dst_doc, text):
    from docx.shared import Pt
    p = dst_doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)


def copy_elements(dst_doc, src_doc, elements):
    body = dst_doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    for el in elements:
        new_el = copy.deepcopy(el)
        for blip in new_el.iter(qn('a:blip')):
            rId = blip.get(qn('r:embed'))
            if not rId:
                continue
            image_part = src_doc.part.related_parts[rId]
            new_rId = dst_doc.part.relate_to(image_part, RT.IMAGE)
            blip.set(qn('r:embed'), new_rId)
        if sect_pr is not None:
            sect_pr.addprevious(new_el)
        else:
            body.append(new_el)


dst = Document()
if src.tables:
    src_style = src.tables[0].style
    try:
        dst.styles.add_style(src_style.name, src_style.type)
    except ValueError:
        pass

append_heading(dst, "PFEM / Transolver Work Summary -- Round 14")
p = dst.add_paragraph()
p.add_run(
    "This round covers Timon's item 4 (design and build a new, harder, more "
    "realistic 3D example). Two candidates were prepared in parallel; Omar's "
    "own detailed 11-point technical review of both was carried out before "
    "this write-up, and this Summary also covers the real GPU result that "
    "followed and an honest account of two genuine fix attempts for a "
    "newly-found QoI limitation. As with previous rounds, everything below "
    "is copied verbatim (same tables, same text) out of the now-updated "
    "cumulative Report (PFEM_Transolver_Report_2026-09-21.docx) -- nothing "
    "summarized or re-derived. Neither candidate's dataset generation or "
    "neural-operator training has started; that step is gated on Timon's "
    "choice between them."
).italic = True

elements = slice_by_anchor_text(
    src,
    "11. Third Benchmark Candidate (in progress): 3D Hyperelastic "
    "Rubber-Mount Bushing vs. Tire Sector",
    "Both candidates are prepared to a real, working, and honestly "
    "characterized state -- B3 with a GPU-confirmed mesh-convergence study "
    "and a full required-resolution table; the tire sector as a genuine, "
    "solving, but deliberately lightweight preliminary alternative. Per "
    "the project's own standing discipline, no dataset generation or "
    "neural-operator training has started for either candidate; that step "
    "is gated on Timon's choice between them.",
)
copy_elements(dst, src, elements)

dst.save(OUT)
print('Saved', OUT)

check = Document(OUT)
print('paragraphs:', len(check.paragraphs))
print('tables:', len(check.tables))
print('images:', len(check.inline_shapes))
