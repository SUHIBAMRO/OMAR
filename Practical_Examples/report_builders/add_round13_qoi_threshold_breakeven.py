"""Round-13 (Timon's newest email, item 1): "for a few QoI-accuracy
thresholds (e.g. 1%/2%/5%), per QoI, the minimum FEM resolution N required
and the operator's own break-even point there."

Data sources (all real, already-verified numbers -- nothing re-derived here):
  - FEM accuracy per N: round12_consistent_field_qoi_<case>.json (the
    already fingerprint-verified, consistent-field data from round 12 --
    FEM and the operator solved on the identical field, both scored
    against one real fine reference, N=201).
  - FEM per-sample GPU-native timing at every LOW_N: new real A100 run,
    2026-09-19 (GPU_FEM_Timing_LowN_AllCases.ipynb, cross-checked against
    the raw Drive JSON for B1xNeo-Hookean before trusting the rest).
  - NO's own compile+TF32 inference cost at N=1401 and each case's real
    training wall-clock: already published (Round12_TrainingMetadata_
    CompileTF32.ipynb, see PROJECT_STATUS.md "Point 3 (compile+TF32...)"
    and apply_round12_points2_3.py's own TRAINING_ROWS).

Omar's own choice (2026-09-19, asked directly): one table per case (6
tables), not one combined summary table -- same per-case pattern as the
existing Tables 18-R10o..z.
"""
import json
import os

from docx import Document
from docx.oxml.ns import qn

DELIV = '/home/user/OMAR/advisor_feedback'
SCRATCH = '/tmp/claude-0/-home-user/64d7c4d8-d5f0-5686-a58f-aa87abfd4ba4/scratchpad/round13_timing'

REPORT_SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-18.docx')
REPORT_DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-19.docx')

CASE_LABELS = {
    'B1_neo_hookean': 'B1 x Neo-Hookean', 'B1_mooney_rivlin': 'B1 x Mooney-Rivlin',
    'B1_arruda_boyce': 'B1 x Arruda-Boyce', 'B2_neo_hookean': 'B2 x Neo-Hookean',
    'B2_mooney_rivlin': 'B2 x Mooney-Rivlin', 'B2_arruda_boyce': 'B2 x Arruda-Boyce',
}
CASE_ORDER = list(CASE_LABELS)
ACC_DIR = '/tmp/claude-0/-home-user/64d7c4d8-d5f0-5686-a58f-aa87abfd4ba4/scratchpad/round12c'
CASE_ACC_FILE = {
    'B1_neo_hookean': 'b1_nh', 'B1_mooney_rivlin': 'b1_mr', 'B1_arruda_boyce': 'b1_ab',
    'B2_neo_hookean': 'b2_nh', 'B2_mooney_rivlin': 'b2_mr', 'B2_arruda_boyce': 'b2_ab',
}

# Real, already-published numbers (see docstring) -- not re-derived here.
NO_COMPILE_TF32_MS = {
    'B1_neo_hookean': 394.09, 'B1_mooney_rivlin': 394.19, 'B1_arruda_boyce': 394.21,
    'B2_neo_hookean': 393.80, 'B2_mooney_rivlin': 394.34, 'B2_arruda_boyce': 394.73,
}
TRAINING_SECONDS = {
    'B1_neo_hookean': 41881.28,  # exact, from metrics_history_multires.json
    'B1_mooney_rivlin': 14.81 * 3600, 'B1_arruda_boyce': 10.83 * 3600,
    'B2_neo_hookean': 10.99 * 3600, 'B2_mooney_rivlin': 1.76 * 3600,
    'B2_arruda_boyce': 2.24 * 3600,  # original run, no retrain for this case
}
QOIS = [
    ('l2_rel', 'Displacement L2'), ('h1_semi_rel', 'H1 semi-norm'),
    ('energy_rel', 'Tangent energy'),
    ('reaction_resultant_rel_err', 'Reaction force (B1 only)'),
    ('cauchy_avg_rel_err', 'Region-Cauchy avg'),
]
THRESHOLDS = [0.01, 0.02, 0.05]  # 1% / 2% / 5%, the advisor's own examples
LOW_N = [3, 4, 5, 6, 9, 11, 13, 17, 21, 25, 29, 33, 37, 41, 45, 49]


def find_required_N(rows_by_N, qoi_key, threshold):
    """First N (ascending) where FEM's own error is <= threshold. FEM
    converges monotonically at every N for every case (established
    2026-09-18), so 'first' == 'minimum sufficient'."""
    for N in LOW_N:
        v = rows_by_N[N]['fem'].get(qoi_key)
        if v is None:
            return None, None
        if v <= threshold:
            return N, v
    return None, None


def compute_results():
    results = {}
    for case_id in CASE_ORDER:
        with open(os.path.join(ACC_DIR, f'round12_consistent_field_qoi_{CASE_ACC_FILE[case_id]}.json')) as f:
            acc = json.load(f)
        rows_by_N = {r['N']: r for r in acc['rows']}
        with open(os.path.join(SCRATCH, f'gpu_fem_timing_lowN_{case_id}.json')) as f:
            timing = json.load(f)
        fem_ms_by_N = {r['N']: r['per_sample_ms'] for r in timing['rows']}

        no_ms = NO_COMPILE_TF32_MS[case_id]
        train_s = TRAINING_SECONDS[case_id]
        case_out = []
        for qoi_key, qoi_label in QOIS:
            if rows_by_N[LOW_N[0]]['fem'].get(qoi_key) is None:
                continue  # e.g. reaction for B2 cases
            for thr in THRESHOLDS:
                N, achieved = find_required_N(rows_by_N, qoi_key, thr)
                if N is None:
                    case_out.append({'qoi': qoi_label, 'threshold_pct': thr * 100,
                                      'required_N': None, 'note': 'not reached by N=49'})
                    continue
                fem_ms = fem_ms_by_N[N]
                saving_ms = fem_ms - no_ms
                if saving_ms > 0:
                    be_note = f'{train_s / (saving_ms / 1000.0):.0f} samples'
                else:
                    be_note = 'FEM already cheaper -- NO never breaks even'
                case_out.append({
                    'qoi': qoi_label, 'threshold_pct': thr * 100, 'required_N': N,
                    'achieved_pct': achieved * 100, 'fem_ms_per_sample': fem_ms, 'note': be_note,
                })
        results[case_id] = case_out
    return results


RESULTS = compute_results()

LETTERS = list('abcdef')
TABLE_ID = {case_id: f'18-R11{LETTERS[i]}' for i, case_id in enumerate(CASE_ORDER)}

HEADER = ['QoI', 'Threshold', 'Required N', 'FEM achieved', 'FEM ms/sample', 'Break-even']


def rows_for(case_id):
    rows = []
    for r in RESULTS[case_id]:
        if r['required_N'] is None:
            rows.append([r['qoi'], f"{r['threshold_pct']:.0f}%", 'n/a', 'n/a', 'n/a', r['note']])
        else:
            rows.append([
                r['qoi'], f"{r['threshold_pct']:.0f}%", str(r['required_N']),
                f"{r['achieved_pct']:.2f}%", f"{r['fem_ms_per_sample']:.1f} ms", r['note'],
            ])
    return rows


def caption(case_id):
    tid = TABLE_ID[case_id]
    label = CASE_LABELS[case_id]
    extra = ''
    if case_id.startswith('B2'):
        extra = ' Reaction force is n/a (no reaction resultant defined for the B2 ring geometry).'
    return (
        f"Table {tid}. {label}: minimum GPU-native-FEM resolution N required to reach each "
        f"QoI-accuracy threshold (FEM's own error, from Table {TABLE_ID_PREV[case_id]}/"
        f"{TABLE_ID_PREV_CAUCHY[case_id]}), that resolution's own real per-sample GPU-FEM cost, "
        f"and the operator's break-even point against it (training wall-clock divided by the "
        f"per-sample saving vs. the operator's own compile+TF32 cost at N=1401, its real "
        f"deployment resolution).{extra}"
    )


TABLE_ID_PREV = {case_id: f'18-R10{c}' for case_id, c in
                  zip(CASE_ORDER, ['o', 'q', 's', 'u', 'w', 'y'])}
TABLE_ID_PREV_CAUCHY = {case_id: f'18-R10{c}' for case_id, c in
                          zip(CASE_ORDER, ['p', 'r', 't', 'v', 'x', 'z'])}

INTRO_TEXT = (
    "Round-13 (the advisor's follow-up email after round 12): finalize the low-N "
    "comparison and add one new analysis -- for a few QoI-accuracy thresholds (1%, 2%, "
    "5%, the advisor's own examples), per QoI, the minimum GPU-native-FEM resolution N "
    "required to reach it, and the operator's own break-even point there. This reuses "
    "round 12's own consistent-field accuracy data (FEM and the operator on the "
    "identical field, scored against one real fine reference) for the accuracy side, "
    "and adds one new real measurement for the cost side: our own GPU-native FEM "
    "solver's (gpu_fem_solver.py) per-sample wall-clock at every one of the sixteen "
    "LOW_N resolutions, not just the single N=11 point the earlier accuracy-matched "
    "break-even table (Table 18-R10e) used. Break-even always uses the operator's own "
    "compile+TF32 cost at N=1401 (its real deployment resolution, essentially "
    "case-independent at ~394ms -- see Table 18-R10h) against that resolution's own "
    "training wall-clock -- the SAME operator baseline as the existing break-even "
    "tables, not a new one. One table per case below."
)

CLOSING_DISCUSSION = (
    "Two findings worth stating plainly. First, across all six cases and every "
    "threshold tested, the required FEM resolution never exceeds N=41, and its "
    "own per-sample cost there (up to ~7.1s) never comes close to the operator's "
    "own ~394ms -- FEM never becomes cheaper than the operator at any accuracy level "
    "tested, so a break-even point always exists and is always finite; the six cases' "
    "break-even sample counts differ almost entirely by each case's own training cost, "
    "not by its accuracy numbers (compare B1xMooney-Rivlin's ~34,500 samples, driven by "
    "its 14.81h training run, against B2xMooney-Rivlin's ~4,000 samples, driven by its "
    "own much cheaper 1.76h run, at the SAME L2@2% threshold). Second, H1 semi-norm and "
    "tangent energy never reach the 1% threshold anywhere in the tested range (N=3..49) "
    "for any of the six cases -- consistent with this project's own repeated finding "
    "that energy/H1-type norms converge more slowly than displacement L2, and a real, "
    "disclosed gap: a 1% energy/H1 target would need FEM resolutions finer than this "
    "sweep covers, not a number this table can currently supply."
)


# =======================================================================
# docx surgery helpers (same pattern as rebuild_round12_point1_consistent_field.py)
# =======================================================================
def body_children(doc):
    return list(doc.element.body)


def elem_text(el):
    if el.tag != qn('w:p'):
        return None
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def find_index(children, text):
    matches = [i for i, el in enumerate(children)
               if elem_text(el) is not None and elem_text(el).strip() == text]
    assert len(matches) == 1, f'{len(matches)} matches for {text!r}'
    return matches[0]


def _elem_of(x):
    if hasattr(x, '_p'):
        return x._p
    if hasattr(x, '_tbl'):
        return x._tbl
    return x


def insert_paragraph_after(doc, anchor, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    p_elem = p._p
    p_elem.getparent().remove(p_elem)
    _elem_of(anchor).addnext(p_elem)
    return p_elem


def insert_table_after(doc, anchor, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    for j, h in enumerate(header):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.rows[i].cells[j].text = v
    tbl_elem = tbl._tbl
    tbl_elem.getparent().remove(tbl_elem)
    _elem_of(anchor).addnext(tbl_elem)
    return tbl_elem


# =======================================================================
# apply
# =======================================================================
doc = Document(REPORT_SRC)
paras = list(doc.paragraphs)
before_p, before_t, before_i = len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes)

anchor_text = [p.text for p in paras if 'Two findings worth stating plainly' in p.text
               and 'true max is markedly noisier' in p.text][0]
children = body_children(doc)
anchor_idx = find_index(children, anchor_text)
anchor = children[anchor_idx]

p = insert_paragraph_after(doc, anchor, INTRO_TEXT)
for case_id in CASE_ORDER:
    p = insert_paragraph_after(doc, p, CASE_LABELS[case_id] + ':', bold=True)
    p = insert_paragraph_after(doc, p, caption(case_id))
    p = insert_table_after(doc, p, HEADER, rows_for(case_id))
p = insert_paragraph_after(doc, p, CLOSING_DISCUSSION)

doc.save(REPORT_DST)
print('Saved', REPORT_DST)

check = Document(REPORT_DST)
after_p, after_t, after_i = len(check.paragraphs), len(check.tables), len(check.inline_shapes)
print(f'paragraphs: {before_p}->{after_p}, tables: {before_t}->{after_t}, images: {before_i}->{after_i}')
expected_new_paras = 1 + 6 * 2 + 1  # intro + (label+caption)*6 + closing
print(f'expected new paragraphs: {expected_new_paras}, expected new tables: 6')
