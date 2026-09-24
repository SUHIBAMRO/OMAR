"""Consistency audit fix, 2026-09-24, per Omar's explicit request to make
sure everything real in the project is reflected in the report, and
matches, before he sends both files onward.

Found: Section 11 (the 3D candidate work: B3 vs. rigid-shim, now a full,
decided result) is INVISIBLE from both of the report's own "what's the
status of everything" locations -- the Executive Summary's own feedback-
tracking table (8 rows, none mentioning the 3D candidate) and Section
10's "The remaining items are:" list (6 items, all predating this work).
This is not a contradiction (nothing in those two places says anything
FALSE about the 3D candidate work), but it is a real completeness gap:
someone reading only the Executive Summary or only Section 10 would not
know Section 11 exists or that a final decision was made. This matches
round 14's own documented, deliberate choice at the time ("started after
that conclusion was written... not a retroactive edit") -- reasonable
while item 4 was still in progress, but no longer complete now that a
real decision exists.

Fix, in place (unlike rounds 15/16, this is NOT a Section-11 rewrite, so
there is no earlier "base" to regenerate from -- this script loads the
already-complete 2026-09-24b.docx and edits two specific, pre-existing
locations directly): (1) a 9th row is added to the Executive Summary's
feedback-tracking table for Timon's item 4, pointing to Section 11; (2) a
7th bullet is inserted into Section 10's "remaining items" list with the
same pointer. Both are one or two sentences, factual, no new claims
beyond what Section 11 itself already states.
"""
import copy
import os

from docx import Document
from docx.oxml.ns import qn

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-24b.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_2026-09-24c.docx')

doc = Document(SRC)

# ------------------------------------------------------------------
# (1) Executive Summary feedback-tracking table: add a 9th row.
# ------------------------------------------------------------------
table = doc.tables[0]
header_cells = [c.text.strip() for c in table.rows[0].cells]
assert header_cells == ['#', 'point', 'Headline result'], header_cells

row = table.add_row()
row.cells[0].text = '9'
row.cells[1].text = ('New, harder, more realistic 3D example (Timon item '
                      '4)')
row.cells[2].text = (
    'Two candidates (B3 with a sharper groove; a laminated rubber-mount '
    'bushing with rigid steel shims) both brought to a full, GPU-'
    'confirmed mesh-convergence result. Prof. Rabczuk selected B3 for '
    'the final FEM-versus-operator comparison, using the 950,400-'
    'element solution as the common reference; the geometry/material/'
    'loading parameters that will vary in the operator training '
    'dataset were also decided. See Section 11.'
)

# ------------------------------------------------------------------
# (2) Section 10 "remaining items" list: insert a 7th bullet, right
# after the existing 6, before the blank paragraph that follows them.
# ------------------------------------------------------------------
paras = doc.paragraphs
last_item_idx = None
for i, p in enumerate(paras):
    if p.style.name == 'List Paragraph' and 'resolution-invariance study is reported' in p.text:
        last_item_idx = i
assert last_item_idx is not None, 'could not find the last "remaining items" bullet'

last_item_para = paras[last_item_idx]
new_p_el = copy.deepcopy(last_item_para._p)
# Clear the copied paragraph's runs (keep its paragraph-level formatting,
# e.g. the List Paragraph bullet numbering/indentation) and write fresh text.
for run_el in new_p_el.findall(qn('w:r')):
    new_p_el.remove(run_el)

last_item_para._p.addnext(new_p_el)

from docx.text.paragraph import Paragraph
new_para = Paragraph(new_p_el, last_item_para._parent)
new_para.add_run(
    'The new, harder, more realistic 3D example requested separately '
    '(item 4) has reached a final candidate decision: see Section 11 '
    'for the full mesh-convergence results of both candidates prepared '
    'and Prof. Rabczuk’s selection of B3. Dataset generation and '
    'operator training for it have not started.'
)

doc.save(DST)
print('Saved', DST)

check = Document(DST)
print('paragraphs:', len(check.paragraphs))
print('tables:', len(check.tables))
print('table[0] rows:', len(check.tables[0].rows))
