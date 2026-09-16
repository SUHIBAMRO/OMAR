"""Companion to add_mms_bodyforce_clarification_to_report.py -- same
clarifying sentence, condensed, added to the Summary's own parallel MMS
section. See that script's own docstring for full rationale."""
import copy
import os

from docx import Document
from docx.text.paragraph import Paragraph

DELIV = '/home/user/OMAR/advisor_feedback'
SRC = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-16d.docx')
DST = os.path.join(DELIV, 'PFEM_Work_Summary_2026-09-16e.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


anchor = find_para_exact(
    "The third leg is now measured. The report's trained operators cannot be "
    "pointed at a manufactured problem no body-force term in the energy "
    "functional, no input channel to carry one so a separate operator was trained "
    "for this study on the same Q4 mesh, minimizing the same discrete potential "
    "energy and scored by the same error routine as the two solvers."
)

insert_after(anchor,
    "Fairness point the advisor raised: since each manufactured solution "
    "generates its own body-force field b = -Div P, how does that varying "
    "forcing reach the operator -- via the family's (alpha, beta) parameters, or "
    "as a direct field? Checked in code: the operator's forcing input and FEM's "
    "own right-hand side are built by the same assembly call with identical "
    "arguments, fed to the operator directly as a per-node channel (only "
    "normalized statistically); (alpha, beta) only generate the family and are "
    "never seen by the operator itself. The operator therefore receives the same "
    "discretized load FEM does, not a lower-dimensional encoding of it.")

doc.save(DST)
print('Saved', DST)
