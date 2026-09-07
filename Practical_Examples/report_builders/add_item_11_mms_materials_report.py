"""Item #11: extend the manufactured-solution verification (Report
section 8.11, Tables 22-24) beyond Neo-Hookean to Mooney-Rivlin and
Arruda-Boyce, into the REAL, already-updated deliverable
(PFEM_Transolver_Report_updated_2026-09-07.docx -- the file with items
#1/#2/#5/#6/#7/#8/#14 already applied).

New tables 22a/22b/23a follow the existing 24a/24b/24c/... convention
of appending letters to an existing table number for closely related
follow-on content, so nothing downstream needs renumbering.
"""
import copy
import json
import os

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Row
from docx.text.paragraph import Paragraph

PF = '/home/user/OMAR/Practical_Examples/omar_pfem'
DELIV = '/tmp/claude-0/-home-user/64d7c4d8-d5f0-5686-a58f-aa87abfd4ba4/scratchpad/deliverables'
SRC = os.path.join(DELIV, 'PFEM_Transolver_Report_updated_2026-09-07.docx')
DST = os.path.join(DELIV, 'PFEM_Transolver_Report_updated_2026-09-07b.docx')

doc = Document(SRC)
ORIGINAL = list(doc.paragraphs)


def find_para_exact(text):
    hits = [p for p in ORIGINAL if p.text.strip() == text]
    assert len(hits) == 1, f'{len(hits)} paragraphs exactly match {text!r}'
    return hits[0]


def insert_after(anchor, text):
    new_p_el = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p_el)
    np = Paragraph(new_p_el, anchor._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def insert_table_after(anchor_para, header, rows):
    """Builds a table matching the document's existing table style and
    inserts it immediately after anchor_para (a Paragraph, not a table --
    tables get their own addnext target below)."""
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = doc.tables[0].style
    pr = doc.tables[0]._tbl.find(qn('w:tblPr'))
    if pr is not None:
        old = tbl._tbl.find(qn('w:tblPr'))
        if old is not None:
            tbl._tbl.remove(old)
        tbl._tbl.insert(0, copy.deepcopy(pr))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    tbl._tbl.getparent().remove(tbl._tbl)
    anchor_para._p.addnext(tbl._tbl)
    return tbl


def insert_after_table(table, style_para, text):
    new_p_el = copy.deepcopy(style_para._p)
    table._tbl.addnext(new_p_el)
    np = Paragraph(new_p_el, style_para._parent)
    for r in list(np.runs):
        r._r.getparent().remove(r._r)
    np.add_run(text)
    return np


def load_rows(material):
    d = json.load(open(os.path.join(PF, 'point9_results', f'mms_B1_{material}.json')))
    return d['rows'], d['convergence_rates'], d['rate_check']


def fmt(v):
    return f'{v:.3e}'


def table_rows_for(rows):
    out = []
    for r in rows:
        n_dof = f"{r['n_dof']:,}" if r['n_dof'] < 1000 else f"{r['n_dof']:,}"
        out.append([r['order'], r['N'], n_dof, fmt(r['L2_rel']), fmt(r['H1_semi_rel']),
                    fmt(r['stress_rel_L2']), fmt(r['energy_rel'])])
    return out


def rate_rows_for(material_label, rates):
    out = []
    for order in ('Q4', 'Q9'):
        for norm_key, norm_label, theory in [
            ('L2', 'L2', 2 if order == 'Q4' else 3),
            ('H1_semi', 'H1 semi-norm', 1 if order == 'Q4' else 2),
            ('stress', 'Stress', 1 if order == 'Q4' else 2),
            ('energy', 'Energy', 2 if order == 'Q4' else 4),
        ]:
            r = rates[order][norm_key]
            pw = ', '.join(f'{v:.2f}' for v in r['pairwise'])
            out.append([material_label, order, norm_label, f"{r['rate']:.2f}", theory, pw])
    return out


mr_rows, mr_rates, mr_check = load_rows('mooney_rivlin')
ab_rows, ab_rates, ab_check = load_rows('arruda_boyce')
assert mr_check == {'Q4': 'as expected', 'Q9': 'as expected'}
assert ab_check == {'Q4': 'as expected', 'Q9': 'as expected'}

HEADER = ['Order', 'N', 'DOF', 'L2', 'H1 semi-norm', 'Stress', 'Energy']

anchor = find_para_exact(
    'The inversion is consistent with what the training principle optimises. The energy '
    'functional is built from the deformation gradient, so the strain and stress fields '
    'are the quantities the loss sees directly and constrains hardest; the displacement '
    'itself is only pinned down through them, anchored by the boundary mask. A '
    'gradient-accurate field that drifts slightly in magnitude is exactly what that '
    'objective tolerates. The same inversion appears in an earlier, independent run at '
    'N = 9 on CPU (1.35× in H1 against 4.71× in L2), so it is not an artefact of one '
    'training run. It is a property worth stating plainly: for a physics-informed '
    'operator, an L2 displacement error overstates how wrong the mechanics are.'
)

p1 = insert_after(
    anchor,
    'Verification across materials. Tables 22-24 verify Neo-Hookean only. The same '
    'derivation (a chosen displacement field, the body force solved backwards from it '
    'by nested autodiff and checked against a central finite difference) and the same '
    'B1 mesh sequence were repeated for the two other constitutive models used '
    'elsewhere in this report, Mooney-Rivlin and Arruda-Boyce, to confirm that the '
    'manufactured-solution check is not accidental to one energy functional. Both use '
    'the identical u* = 0.05*(sin(pi x)sin(pi y), 0.7*sin(pi x)sin(pi y)) field and the '
    'same uniform E = 1000, nu = 0.3 (plane strain) material field as Table 22, so the '
    'three materials differ only in psi(F) and its derived body force -- not in mesh, '
    'boundary conditions, or the field being reproduced.'
)

t22a = insert_table_after(p1, HEADER, table_rows_for(mr_rows))
p2 = insert_after_table(
    t22a, p1,
    'Table 22a. Q4 and Q9 against the manufactured solution, B1 geometry, '
    'Mooney-Rivlin, FP64. Same columns and error definitions as Table 22, one '
    'additional resolution (N = 33) since the cost of one more point at this material '
    'was small.'
)

t22b = insert_table_after(p2, HEADER, table_rows_for(ab_rows))
p3 = insert_after_table(
    t22b, p2,
    'Table 22b. Q4 and Q9 against the manufactured solution, B1 geometry, '
    'Arruda-Boyce, FP64. Same columns, error definitions, and resolutions as Table 22a.'
)

p4 = insert_after(
    p3,
    'The raw error values differ slightly from Table 22\'s Neo-Hookean numbers at '
    'matched N and order -- expected, since a different psi(F) gives a different '
    'tangent stiffness and therefore a different discretization error for the same '
    'exact displacement field -- but the pattern is the same: Q9 wins at equal DOF and '
    'the gap widens with refinement, exactly as in Table 22.'
)

rate_header = ['Material', 'Order', 'Norm', 'Observed rate', 'Theory', 'Pairwise']
rate_rows = rate_rows_for('Mooney-Rivlin', mr_rates) + rate_rows_for('Arruda-Boyce', ab_rates)
t23a = insert_table_after(p4, rate_header, rate_rows)
p5 = insert_after_table(
    t23a, p4,
    'Table 23a. Observed convergence rates for Mooney-Rivlin and Arruda-Boyce, fitted '
    'the same way as Table 23 (least squares on log h, three consecutive-pair points '
    'from the four resolutions in Tables 22a/22b). All sixteen rows land on their '
    'theoretical rate, matching Table 23\'s own Neo-Hookean result -- the verification '
    'holds for every constitutive model this report trains an operator on, not only '
    'the one exercised in Tables 22-24.'
)

doc.save(DST)
print('Saved', DST)
